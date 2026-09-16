"""
Markdown to Word converter.
"""
from datetime import datetime, timezone
import getpass
import os
from pathlib import Path
import re
import shutil
import tempfile
from typing import Any, Dict, Iterable, Optional, Union
from zipfile import ZIP_DEFLATED, ZipFile
import xml.etree.ElementTree as ET

from docx import Document
from docx.oxml.ns import qn

from md2docx.config_utils import clone_config, merge_config
from md2docx.styles import StyleManager
from md2docx.parser import MarkdownParser
from md2docx.mermaid_converter import MermaidReport
from md2docx.omml_converter import FormulaReport


EXTENDED_PROPERTIES_NS = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"


class Converter:
    """Main converter class for Markdown to Word conversion."""

    @property
    def mermaid_report(self) -> MermaidReport:
        """Return Mermaid outcomes for the most recent document conversion."""
        if self.parser.renderer is None:
            return MermaidReport()
        return self.parser.renderer.mermaid_report

    @property
    def formula_report(self) -> FormulaReport:
        """Return native, fallback, and failed formula outcomes for the last document."""
        if self.parser.renderer is None:
            return FormulaReport()
        return self.parser.renderer.formula_report
    
    def __init__(
        self,
        template: Optional[str] = None,
        style_config: Optional[str] = None,
        config_override: Optional[Dict[str, Any]] = None,
        config_data: Optional[Dict[str, Any]] = None,
        word_template: Optional[str] = None,
    ) -> None:
        """
        Initialize converter.
        
        Args:
            template: Name of predefined style template
            style_config: Path to custom YAML style configuration
            config_override: Runtime config overrides merged over the loaded style config
            config_data: Complete in-memory style config that bypasses disk defaults
            word_template: Path to a .docx template whose package-owned headers and
                footers should be preserved in the generated document
        """
        self.template = template
        self.style_config = style_config
        self.config_override = config_override or {}
        self.config_data = clone_config(config_data) if config_data is not None else None
        self.word_template = word_template
        
        # Initialize style manager
        config_path = style_config or template
        if self.config_data is None:
            self.style_manager = StyleManager(config_path)
        else:
            self.style_manager = StyleManager.from_config(self.config_data)
        if self.config_override:
            self.style_manager.config = merge_config(
                self.style_manager.config,
                self.config_override,
            )
        
        # Initialize parser
        self.parser = MarkdownParser(self.style_manager)
    
    def convert(self, md_path: str, docx_path: str) -> None:
        """
        Convert Markdown file to Word document.
        
        Args:
            md_path: Path to input Markdown file
            docx_path: Path to output Word document
            
        Raises:
            FileNotFoundError: If input file doesn't exist
            IOError: If there's an error reading/writing files
        """
        md_file = Path(md_path)
        
        if not md_file.exists():
            raise FileNotFoundError(f"Markdown file not found: {md_path}")
        
        # Read Markdown content
        with open(md_file, 'r', encoding='utf-8') as f:
            md_content = f.read()
        
        # Convert to Word
        self.convert_string(md_content, docx_path, base_dir=md_file.parent)
    
    def convert_string(
        self,
        md_content: str,
        docx_path: str,
        base_dir: Optional[Union[str, Path]] = None,
    ) -> None:
        """
        Convert Markdown string to Word document.
        
        Args:
            md_content: Markdown content as string
            docx_path: Path to output Word document
            base_dir: Directory used to resolve relative image paths
            
        Raises:
            IOError: If there's an error writing the file
        """
        # Create document
        doc = self.to_document(md_content, base_dir=base_dir)
        
        # Save document
        output_path = Path(docx_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        
        doc.save(str(output_path))
        self._update_extended_properties(output_path, doc)
    
    def to_document(
        self,
        md_content: str,
        base_dir: Optional[Union[str, Path]] = None,
    ) -> Document:
        """
        Convert Markdown string to Document object.
        
        Args:
            md_content: Markdown content as string
            base_dir: Directory used to resolve relative image paths
            
        Returns:
            python-docx Document object
        """
        # Apply document-level styles
        doc_style = self.style_manager.get_document_style()

        if self.word_template:
            doc = self._load_word_template(self.word_template)
        else:
            doc = Document()
            self._apply_document_settings(doc, doc_style)
        
        # Parse Markdown and add content to document
        self.parser.parse(md_content, doc, base_dir=base_dir)
        self._apply_core_properties(doc, md_content)
        
        return doc

    def _load_word_template(self, template_path: str) -> Document:
        """Load a single-section DOCX template and clear only its body content."""
        path = Path(template_path).expanduser()
        if not path.exists() or not path.is_file():
            raise FileNotFoundError(f"Word template not found: {path}")
        if path.suffix.lower() != ".docx":
            raise ValueError(f"Word template must be a .docx file: {path}")

        doc = Document(str(path))
        if len(doc.sections) != 1:
            raise ValueError("Word template must contain exactly one section")

        body = doc._element.body
        for child in list(body):
            if child.tag != qn("w:sectPr"):
                body.remove(child)
        return doc

    def _apply_core_properties(self, doc: Document, md_content: str) -> None:
        """
        Replace python-docx default metadata with document-specific values.
        """
        metadata = self.style_manager.get_style("metadata", {})
        if not isinstance(metadata, dict):
            metadata = {}

        author = str(
            metadata.get("author")
            or getpass.getuser()
            or os.environ.get("USER")
            or os.environ.get("USERNAME")
            or "md2docx"
        )
        now = datetime.now(timezone.utc)
        props = doc.core_properties

        props.author = author
        props.last_modified_by = str(metadata.get("last_modified_by") or author)
        props.title = str(
            metadata.get("title")
            or self._extract_title(md_content)
            or self._first_document_text(doc)
            or ""
        )
        props.subject = str(metadata.get("subject") or "")
        props.keywords = self._format_keywords(metadata.get("keywords"))
        props.comments = str(metadata.get("comments") or "")
        props.category = str(metadata.get("category") or "")
        props.created = now
        props.modified = now

    def _extract_title(self, md_content: str) -> str:
        """Return the first Markdown heading as a document title."""
        for line in md_content.splitlines():
            match = re.match(r"^\s{0,3}#{1,6}\s+(.+?)\s*#*\s*$", line)
            if match:
                return self._clean_metadata_text(match.group(1))
        return ""

    def _first_document_text(self, doc: Document) -> str:
        """Return the first non-empty paragraph/table text from a document."""
        for text in self._iter_document_text(doc):
            cleaned = text.strip()
            if cleaned:
                return self._clean_metadata_text(cleaned)
        return ""

    def _format_keywords(self, keywords: Any) -> str:
        """Format keywords from configuration into a metadata string."""
        if isinstance(keywords, (list, tuple, set)):
            return ", ".join(str(keyword) for keyword in keywords if str(keyword).strip())
        return str(keywords or "")

    def _clean_metadata_text(self, text: str) -> str:
        """Remove lightweight Markdown markers from metadata values."""
        text = re.sub(r"!\[([^\]]*)\]\([^)]+\)", r"\1", text)
        text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
        text = re.sub(r"[*_`~]+", "", text)
        return text.strip()

    def _update_extended_properties(self, docx_path: Path, doc: Document) -> None:
        """
        Refresh docProps/app.xml statistics after python-docx saves the package.
        """
        stats = self._collect_document_stats(doc)
        replacement_path = None
        try:
            with tempfile.NamedTemporaryFile(
                suffix=".docx",
                delete=False,
                dir=str(docx_path.parent),
            ) as tmp_file:
                replacement_path = Path(tmp_file.name)

            with ZipFile(docx_path, "r") as source, ZipFile(
                replacement_path,
                "w",
                compression=ZIP_DEFLATED,
            ) as target:
                for item in source.infolist():
                    data = source.read(item.filename)
                    if item.filename == "docProps/app.xml":
                        data = self._rewrite_app_properties(data, stats)
                    target.writestr(item, data)

            shutil.move(str(replacement_path), str(docx_path))
        finally:
            if replacement_path is not None and replacement_path.exists():
                replacement_path.unlink(missing_ok=True)

    def _rewrite_app_properties(self, app_xml: bytes, stats: Dict[str, int]) -> bytes:
        """Update extended document statistics in app.xml."""
        ET.register_namespace("", EXTENDED_PROPERTIES_NS)
        root = ET.fromstring(app_xml)

        self._set_extended_property(root, "Pages", max(1, stats["pages"]))
        self._set_extended_property(root, "Words", stats["words"])
        self._set_extended_property(root, "Characters", stats["characters"])
        self._set_extended_property(root, "CharactersWithSpaces", stats["characters_with_spaces"])
        self._set_extended_property(root, "Lines", stats["lines"])
        self._set_extended_property(root, "Paragraphs", stats["paragraphs"])
        self._set_extended_property(root, "TotalTime", 1)
        self._set_extended_property(root, "AppVersion", "16.0000")

        return ET.tostring(root, encoding="utf-8", xml_declaration=True)

    def _set_extended_property(self, root: ET.Element, name: str, value: Any) -> None:
        """Set one property in docProps/app.xml, creating it if needed."""
        element = root.find(f"{{{EXTENDED_PROPERTIES_NS}}}{name}")
        if element is None:
            element = ET.SubElement(root, f"{{{EXTENDED_PROPERTIES_NS}}}{name}")
        element.text = str(value)

    def _collect_document_stats(self, doc: Document) -> Dict[str, int]:
        """Collect conservative document statistics for docProps/app.xml."""
        text_parts = [text.strip() for text in self._iter_document_text(doc) if text.strip()]
        full_text = "\n".join(text_parts)
        characters_with_spaces = len(full_text)
        characters = len(re.sub(r"\s+", "", full_text))
        words = len(re.findall(r"[A-Za-z0-9]+|[\u4e00-\u9fff]", full_text))
        paragraphs = len(text_parts)

        return {
            "pages": 1,
            "words": max(1, words) if full_text else 0,
            "characters": characters,
            "characters_with_spaces": characters_with_spaces,
            "lines": max(1, paragraphs) if full_text else 0,
            "paragraphs": paragraphs,
        }

    def _iter_document_text(self, doc: Document) -> Iterable[str]:
        """Yield paragraph and table cell text from a python-docx document."""
        for paragraph in doc.paragraphs:
            yield paragraph.text

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield cell.text
    
    def _apply_document_settings(self, doc: Document, style: dict) -> None:
        """
        Apply document-level settings like page size and margins.
        
        Args:
            doc: python-docx Document object
            style: Document style configuration
        """
        from docx.shared import Inches, Cm, Pt
        from docx.enum.section import WD_SECTION
        
        # Get the default section
        section = doc.sections[0]
        
        # Apply page size
        if 'page_size' in style:
            page_size = style['page_size'].upper()
            if page_size == 'A4':
                section.page_height = Cm(29.7)
                section.page_width = Cm(21.0)
            elif page_size == 'A3':
                section.page_height = Cm(42.0)
                section.page_width = Cm(29.7)
            elif page_size == 'LETTER':
                section.page_height = Inches(11)
                section.page_width = Inches(8.5)
        
        # Apply margins
        if 'margin_top' in style:
            section.top_margin = self._parse_length(style['margin_top'])
        
        if 'margin_bottom' in style:
            section.bottom_margin = self._parse_length(style['margin_bottom'])
        
        if 'margin_left' in style:
            section.left_margin = self._parse_length(style['margin_left'])
        
        if 'margin_right' in style:
            section.right_margin = self._parse_length(style['margin_right'])
    
    def _parse_length(self, length_str: str):
        """
        Parse length string to python-docx length object.
        
        Args:
            length_str: Length string like "2.54cm" or "1in"
            
        Returns:
            Length object (Inches, Cm, Pt, etc.)
        """
        from docx.shared import Inches, Cm, Pt, Mm
        
        length_str = str(length_str).lower().strip()
        
        if length_str.endswith('cm'):
            value = float(length_str[:-2])
            return Cm(value)
        elif length_str.endswith('in'):
            value = float(length_str[:-2])
            return Inches(value)
        elif length_str.endswith('mm'):
            value = float(length_str[:-2])
            return Mm(value)
        elif length_str.endswith('pt'):
            value = float(length_str[:-2])
            return Pt(value)
        else:
            # Default to cm
            return Cm(float(length_str))
