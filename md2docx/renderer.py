"""
Word document renderer module.
"""
from typing import Any, Dict, List, Optional, Tuple, Union
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm, Mm, Emu
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import re
import mistune
from io import BytesIO
from urllib.parse import unquote


TABLE_LAYOUT_ACCENT_GRID = "accent_grid"
TABLE_LAYOUT_THREE_LINE = "three_line"
TABLE_LAYOUT_PLAIN_GRID = "plain_grid"
TABLE_LAYOUTS = frozenset(
    {
        TABLE_LAYOUT_ACCENT_GRID,
        TABLE_LAYOUT_THREE_LINE,
        TABLE_LAYOUT_PLAIN_GRID,
    }
)


class DocxRenderer(mistune.BaseRenderer):
    """Custom mistune renderer that outputs to python-docx."""
    
    def __init__(
        self,
        doc: Document,
        style_manager: Any,
        base_dir: Optional[Union[str, Path]] = None,
    ) -> None:
        """
        Initialize renderer.
        
        Args:
            doc: python-docx Document object
            style_manager: StyleManager instance for style configuration
            base_dir: Directory used to resolve relative Markdown image paths
        """
        super().__init__()
        self.doc = doc
        self.styles = style_manager
        self.base_dir = Path(base_dir).resolve() if base_dir else None
        self._current_paragraph = None
        self.math_formulas = {'inline': [], 'block': []}  # Will be set by parser
        self.outline_processing = {'mode': 'auto', 'active': False}
        
        # Initialize math converter for LaTeX formulas
        from md2docx.math_converter import MathConverter
        from md2docx.omml_converter import OmmlConverter
        self.math_converter = MathConverter(dpi=300)
        self.omml_converter = OmmlConverter()

        # Initialize Mermaid converter for fenced mermaid blocks
        from md2docx.mermaid_converter import MermaidConverter
        mermaid_style = self.styles.get_mermaid_style()
        self.mermaid_converter = MermaidConverter(
            command=mermaid_style.get('command', 'mmdc'),
            output_format=mermaid_style.get('format', 'png'),
            theme=mermaid_style.get('theme'),
            theme_variables=mermaid_style.get('theme_variables'),
            background_color=mermaid_style.get('background_color', 'white'),
            scale=mermaid_style.get('scale'),
        )
        self._ordered_abstract_num_id = self._find_list_abstract_num_id(
            style_id='ListNumber',
            num_format='decimal',
            fallback=7,
        )
    
    def _parse_font_size(self, size_str: str) -> int:
        """
        Parse font size string to points.
        
        Args:
            size_str: Font size string like "12pt" or "14"
            
        Returns:
            Font size in points
        """
        if isinstance(size_str, int):
            return size_str
        
        size_str = str(size_str).lower().strip()
        
        if size_str.endswith('pt'):
            return int(size_str[:-2])
        
        return int(size_str)
    
    def _parse_color(self, color_str: str) -> RGBColor:
        """
        Parse color string to RGBColor.
        
        Args:
            color_str: Color string like "#FF0000" or "FF0000" or "#D14" (short format)
            
        Returns:
            RGBColor object
        """
        color_str = color_str.strip()
        
        if color_str.startswith('#'):
            color_str = color_str[1:]
        
        # Support short format (#RGB -> #RRGGBB)
        if len(color_str) == 3:
            color_str = ''.join([c*2 for c in color_str])
        
        r = int(color_str[0:2], 16)
        g = int(color_str[2:4], 16)
        b = int(color_str[4:6], 16)
        
        return RGBColor(r, g, b)
    
    def _get_alignment(self, alignment_str: str) -> WD_PARAGRAPH_ALIGNMENT:
        """Get alignment enum from string."""
        alignment_map = {
            'left': WD_PARAGRAPH_ALIGNMENT.LEFT,
            'center': WD_PARAGRAPH_ALIGNMENT.CENTER,
            'right': WD_PARAGRAPH_ALIGNMENT.RIGHT,
            'justify': WD_PARAGRAPH_ALIGNMENT.JUSTIFY,
        }
        return alignment_map.get(alignment_str.lower(), WD_PARAGRAPH_ALIGNMENT.LEFT)

    def _get_vertical_alignment(self, alignment_str: str) -> WD_CELL_VERTICAL_ALIGNMENT:
        """Get cell vertical alignment enum from string."""
        alignment_map = {
            'top': WD_CELL_VERTICAL_ALIGNMENT.TOP,
            'center': WD_CELL_VERTICAL_ALIGNMENT.CENTER,
            'bottom': WD_CELL_VERTICAL_ALIGNMENT.BOTTOM,
        }
        return alignment_map.get(alignment_str.lower(), WD_CELL_VERTICAL_ALIGNMENT.TOP)

    def _parse_length(self, length_value: Any, default_unit: str = 'pt'):
        """
        Parse a length value into a python-docx length object.

        Args:
            length_value: Length like "5.5in", "12pt", or a numeric value
            default_unit: Unit to assume for numeric values

        Returns:
            Parsed python-docx length object
        """
        if hasattr(length_value, 'emu'):
            return length_value

        if isinstance(length_value, (int, float)):
            value = float(length_value)
            if default_unit == 'in':
                return Inches(value)
            if default_unit == 'cm':
                return Cm(value)
            if default_unit == 'mm':
                return Mm(value)
            return Pt(value)

        length_str = str(length_value).lower().strip()

        if length_str.endswith('cm'):
            return Cm(float(length_str[:-2]))
        if length_str.endswith('in'):
            return Inches(float(length_str[:-2]))
        if length_str.endswith('mm'):
            return Mm(float(length_str[:-2]))
        if length_str.endswith('pt'):
            return Pt(float(length_str[:-2]))

        if default_unit == 'in':
            return Inches(float(length_str))
        if default_unit == 'cm':
            return Cm(float(length_str))
        if default_unit == 'mm':
            return Mm(float(length_str))
        return Pt(float(length_str))

    def _get_available_page_width(self):
        """Get the usable document width after subtracting page margins."""
        section = self.doc.sections[-1]
        return section.page_width - section.left_margin - section.right_margin

    def _weighted_text_length(self, text: str) -> float:
        """Estimate text width, counting CJK characters as wider than ASCII."""
        weight = 0.0
        for char in text:
            if char.isspace():
                weight += 0.25
            elif '\u4e00' <= char <= '\u9fff':
                weight += 1.8
            elif ord(char) > 127:
                weight += 1.4
            else:
                weight += 1.0
        return weight

    def _plain_token_text(self, token: Dict[str, Any]) -> str:
        """Extract readable text from a mistune token tree without rendering."""
        if 'raw' in token:
            return str(token['raw'])

        if 'text' in token:
            return str(token['text'])

        children = token.get('children', [])
        if isinstance(children, list):
            return ''.join(self._plain_token_text(child) for child in children)

        return ''

    def _collect_table_column_weights(
        self,
        table_head: Optional[Dict[str, Any]],
        table_body: Optional[Dict[str, Any]],
        col_count: int,
    ) -> List[float]:
        """Estimate content weight for each table column."""
        weights = [1.0 for _ in range(col_count)]

        rows = []
        if table_head:
            rows.append(table_head)
        if table_body:
            rows.extend(table_body.get('children', []))

        for row_token in rows:
            cells = row_token.get('children', [])
            for col_idx, cell_token in enumerate(cells[:col_count]):
                text = self._plain_token_text(cell_token).strip()
                weights[col_idx] = max(weights[col_idx], self._weighted_text_length(text))

        return weights

    def _normalize_width_ratios(self, ratios: List[float], minimum: float, maximum: float) -> List[float]:
        """Clamp and normalize width ratios until they sum to 1.0."""
        if not ratios:
            return []

        clamped = [min(max(ratio, minimum), maximum) for ratio in ratios]
        total = sum(clamped)
        if total <= 0:
            return [1 / len(ratios) for _ in ratios]

        normalized = [ratio / total for ratio in clamped]
        for _ in range(4):
            adjusted = [min(max(ratio, minimum), maximum) for ratio in normalized]
            total = sum(adjusted)
            normalized = [ratio / total for ratio in adjusted]

        return normalized

    def _calculate_table_column_widths(
        self,
        table_head: Optional[Dict[str, Any]],
        table_body: Optional[Dict[str, Any]],
        col_count: int,
        table_style: Dict[str, Any],
    ) -> List[int]:
        """Calculate table column widths in EMUs."""
        if col_count <= 0:
            return []

        available_width = int(self._get_available_page_width())
        equal_ratio = 1 / col_count
        strategy = str(table_style.get('column_width_strategy', 'content-weighted')).lower()
        weights = self._collect_table_column_weights(table_head, table_body, col_count)
        weight_total = sum(weights)

        if weight_total <= 0:
            ratios = [equal_ratio for _ in range(col_count)]
        else:
            ratios = [weight / weight_total for weight in weights]

        if strategy == 'balanced':
            ratios = [(ratio * 0.45) + (equal_ratio * 0.55) for ratio in ratios]
            ratios = self._normalize_width_ratios(
                ratios,
                minimum=max(0.12, equal_ratio * 0.72),
                maximum=min(0.48, equal_ratio * 1.45),
            )
        else:
            ratios = self._normalize_width_ratios(
                ratios,
                minimum=max(0.08, equal_ratio * 0.38),
                maximum=min(0.62, equal_ratio * 2.15),
            )

        widths = [max(1, int(round(available_width * ratio))) for ratio in ratios]
        width_delta = available_width - sum(widths)
        if widths:
            widths[-1] += width_delta

        return widths

    def _set_table_column_widths(self, table: Any, widths: List[int]) -> None:
        """Apply fixed table and column widths to a python-docx table."""
        if not widths:
            return

        table.autofit = False
        table.allow_autofit = False

        tbl_pr = table._tbl.tblPr
        tbl_layout = tbl_pr.first_child_found_in('w:tblLayout')
        if tbl_layout is None:
            tbl_layout = OxmlElement('w:tblLayout')
            tbl_pr.append(tbl_layout)
        tbl_layout.set(qn('w:type'), 'fixed')

        table_width = sum(widths)
        tbl_w = tbl_pr.first_child_found_in('w:tblW')
        if tbl_w is None:
            tbl_w = OxmlElement('w:tblW')
            tbl_pr.append(tbl_w)
        tbl_w.set(qn('w:type'), 'dxa')
        tbl_w.set(qn('w:w'), str(int(round(table_width / 635))))

        grid = table._tbl.tblGrid
        for col_idx, width in enumerate(widths):
            if col_idx < len(grid.gridCol_lst):
                grid_col = grid.gridCol_lst[col_idx]
            else:
                grid_col = OxmlElement('w:gridCol')
                grid.append(grid_col)
            grid_col.set(qn('w:w'), str(int(round(width / 635))))

        for row in table.rows:
            for col_idx, cell in enumerate(row.cells):
                if col_idx >= len(widths):
                    continue

                cell.width = Emu(widths[col_idx])
                tc_pr = cell._tc.get_or_add_tcPr()
                tc_w = tc_pr.tcW
                if tc_w is None:
                    tc_w = OxmlElement('w:tcW')
                    tc_pr.append(tc_w)
                tc_w.set(qn('w:type'), 'dxa')
                tc_w.set(qn('w:w'), str(int(round(widths[col_idx] / 635))))

    def _set_cell_margins(
        self,
        cell: Any,
        *,
        top: Optional[Any] = None,
        bottom: Optional[Any] = None,
        left: Optional[Any] = None,
        right: Optional[Any] = None,
    ) -> None:
        """Set explicit Word table cell margins."""
        margins = {
            'top': top,
            'bottom': bottom,
            'left': left,
            'right': right,
        }
        if all(value is None for value in margins.values()):
            return

        tc_pr = cell._tc.get_or_add_tcPr()
        tc_mar = tc_pr.first_child_found_in('w:tcMar')
        if tc_mar is None:
            tc_mar = OxmlElement('w:tcMar')
            tc_pr.append(tc_mar)

        for side, value in margins.items():
            if value is None:
                continue
            margin = tc_mar.find(qn(f'w:{side}'))
            if margin is None:
                margin = OxmlElement(f'w:{side}')
                tc_mar.append(margin)
            margin.set(qn('w:w'), str(int(round(self._parse_length(value) / 635))))
            margin.set(qn('w:type'), 'dxa')

    def _apply_table_cell_margins(self, cell: Any, table_style: Dict[str, Any]) -> None:
        """Apply configured cell margins while preserving default horizontal padding."""
        vertical_margin = table_style.get('cell_margin_vertical', '3pt')
        horizontal_margin = table_style.get('cell_margin_horizontal', '5.4pt')
        self._set_cell_margins(
            cell,
            top=vertical_margin,
            bottom=vertical_margin,
            left=horizontal_margin,
            right=horizontal_margin,
        )

    def _resolve_table_layout(self, table_style: Dict[str, Any]) -> str:
        """Return a supported layout, preserving the historic default."""
        layout = str(
            table_style.get("layout", TABLE_LAYOUT_ACCENT_GRID)
        ).strip().lower()
        if layout not in TABLE_LAYOUTS:
            return TABLE_LAYOUT_ACCENT_GRID
        return layout

    def _set_ooxml_borders(
        self,
        properties: Any,
        container_tag: str,
        borders: Dict[str, Tuple[str, Optional[int]]],
    ) -> None:
        """Set deterministic black borders on table or cell properties."""
        container = properties.first_child_found_in(container_tag)
        if container is None:
            container = OxmlElement(container_tag)
            properties.append(container)

        for side, (value, size) in borders.items():
            border = container.find(qn(f"w:{side}"))
            if border is None:
                border = OxmlElement(f"w:{side}")
                container.append(border)
            border.set(qn("w:val"), value)
            if value == "single" and size is not None:
                border.set(qn("w:sz"), str(size))
                border.set(qn("w:color"), "000000")
            else:
                border.attrib.pop(qn("w:sz"), None)
                border.attrib.pop(qn("w:color"), None)

    def _apply_table_layout(
        self,
        table: Any,
        table_style: Dict[str, Any],
        layout: str,
    ) -> None:
        """Apply only the visual properties owned by a table preset."""
        if layout == TABLE_LAYOUT_ACCENT_GRID:
            style_name = table_style.get("style")
            if style_name:
                try:
                    table.style = style_name
                except KeyError:
                    pass
            return

        table_properties = table._tbl.tblPr
        table_style_element = table_properties.find(qn("w:tblStyle"))
        if table_style_element is not None:
            table_properties.remove(table_style_element)

        if layout == TABLE_LAYOUT_PLAIN_GRID:
            self._set_ooxml_borders(
                table_properties,
                "w:tblBorders",
                {
                    side: ("single", 4)
                    for side in (
                        "top",
                        "left",
                        "bottom",
                        "right",
                        "insideH",
                        "insideV",
                    )
                },
            )
            return

        self._set_ooxml_borders(
            table_properties,
            "w:tblBorders",
            {
                "top": ("single", 12),
                "left": ("nil", None),
                "bottom": ("single", 12),
                "right": ("nil", None),
                "insideH": ("nil", None),
                "insideV": ("nil", None),
            },
        )
        for cell in table.rows[0].cells:
            self._set_ooxml_borders(
                cell._tc.get_or_add_tcPr(),
                "w:tcBorders",
                {"bottom": ("single", 6)},
            )

    def _set_row_repeats_as_table_header(self, row: Any) -> None:
        """Mark a table row to repeat as the header row on each Word page."""
        tr_pr = row._tr.get_or_add_trPr()
        tbl_header = tr_pr.find(qn('w:tblHeader'))
        if tbl_header is None:
            tbl_header = OxmlElement('w:tblHeader')
            tr_pr.append(tbl_header)
        tbl_header.set(qn('w:val'), 'true')

    def _get_available_page_height(self):
        """Get the usable document height after subtracting page margins."""
        section = self.doc.sections[-1]
        return section.page_height - section.top_margin - section.bottom_margin

    def _find_list_abstract_num_id(self, style_id: str, num_format: str, fallback: int) -> int:
        """
        Find the numbering template used by a built-in Word list style.

        We clone the built-in abstract numbering definition and create a fresh
        concrete numbering instance for each Markdown ordered list block. This
        lets every block restart from 1 instead of continuing the previous one.
        """
        numbering = self.doc.part.numbering_part.numbering_definitions._numbering

        matches = numbering.xpath(
            f'./w:abstractNum[w:lvl/w:pStyle[@w:val="{style_id}"]]/@w:abstractNumId'
        )
        if matches:
            return int(matches[0])

        matches = numbering.xpath(
            f'./w:abstractNum[w:lvl/w:numFmt[@w:val="{num_format}"]]/@w:abstractNumId'
        )
        if matches:
            return int(matches[0])

        return fallback

    def _create_ordered_list_num_id(self) -> int:
        """
        Create a fresh numbering sequence for a single Markdown ordered list block.
        """
        numbering = self.doc.part.numbering_part.numbering_definitions._numbering
        num = numbering.add_num(self._ordered_abstract_num_id)
        num.add_lvlOverride(ilvl=0).add_startOverride(1)
        return int(num.numId)

    def _set_paragraph_numbering(self, paragraph: Any, num_id: int, ilvl: int = 0) -> None:
        """
        Attach an explicit numbering sequence to a paragraph.

        We intentionally keep `ilvl=0` because the built-in numbering templates
        shipped with python-docx are single-level. Nested indentation is handled
        separately via paragraph indentation.
        """
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        pPr = paragraph._element.get_or_add_pPr()
        existing_numPr = pPr.find(qn('w:numPr'))
        if existing_numPr is not None:
            pPr.remove(existing_numPr)

        numPr = OxmlElement('w:numPr')

        ilvl_elm = OxmlElement('w:ilvl')
        ilvl_elm.set(qn('w:val'), str(ilvl))
        numPr.append(ilvl_elm)

        numId_elm = OxmlElement('w:numId')
        numId_elm.set(qn('w:val'), str(num_id))
        numPr.append(numId_elm)

        pPr.append(numPr)

    def _outline_mode_active(self) -> bool:
        """Return True when Chinese outline heading handling is active."""
        return bool(self.outline_processing.get('active', False))

    def _match_outline_level(self, text: str) -> Optional[int]:
        """
        Match Chinese document outline markers at the start of a paragraph.

        Supported levels:
        1. `一、`
        2. `（一）`
        3. `1.`
        4. `（1）`
        """
        stripped = text.strip()
        patterns = (
            (1, rf'^[零〇一二三四五六七八九十百千万]+、\s*\S'),
            (2, rf'^[（(][零〇一二三四五六七八九十百千万]+[）)]\s*\S'),
            (3, r'^\d+\.\s*\S'),
            (4, r'^[（(]\d+[）)]\s*\S'),
        )

        for level, pattern in patterns:
            if re.match(pattern, stripped):
                return level

        return None

    def _resolve_style_reference(self, style_name: str) -> Dict[str, Any]:
        """Resolve a configured style name to a concrete style dictionary."""
        if style_name == 'paragraph':
            return self.styles.get_paragraph_style()

        if style_name.startswith('heading'):
            try:
                level = int(style_name.replace('heading', ''))
                return self.styles.get_heading_style(level)
            except ValueError:
                pass

        return self.styles.get_style(style_name, self.styles.get_paragraph_style())

    def _get_outline_style(self, level: int) -> Dict[str, Any]:
        """Get the style used for a detected outline heading level."""
        outline_config = self.styles.get_style('outline', {})
        default_style_map = {
            1: 'heading2',
            2: 'heading3',
            3: 'heading4',
            4: 'paragraph',
        }
        style_name = outline_config.get(f'level{level}_style', default_style_map[level])
        style = dict(self._resolve_style_reference(style_name))

        overrides = outline_config.get(f'level{level}', {})
        if isinstance(overrides, dict):
            style.update(overrides)

        return style

    def _apply_text_paragraph_format(self, paragraph: Any, style: Dict[str, Any]) -> None:
        """Apply alignment, spacing, and indentation to a text paragraph."""
        if 'alignment' in style:
            paragraph.alignment = self._get_alignment(style['alignment'])

        if 'line_spacing' in style:
            paragraph.paragraph_format.line_spacing = style['line_spacing']

        if 'first_line_indent' in style and style['first_line_indent'] > 0:
            char_count = style['first_line_indent']
            font_size_pt = self._parse_font_size(style.get('font_size', '12pt'))
            paragraph.paragraph_format.first_line_indent = Pt(char_count * font_size_pt)

        if 'space_before' in style:
            paragraph.paragraph_format.space_before = Pt(self._parse_font_size(style['space_before']))

        if 'space_after' in style:
            paragraph.paragraph_format.space_after = Pt(self._parse_font_size(style['space_after']))

    def _render_outline_paragraph(self, text: str, level: int) -> None:
        """Render a detected outline heading as a styled paragraph."""
        style = self._get_outline_style(level)
        paragraph = self.doc.add_paragraph()
        self._add_formatted_text(paragraph, text, style)

        if 'bold' in style or 'italic' in style:
            for run in paragraph.runs:
                if 'bold' in style:
                    run.font.bold = style['bold']
                if 'italic' in style:
                    run.font.italic = style['italic']

        self._apply_text_paragraph_format(paragraph, style)

    def _get_image_pixel_size(self, img_bytes: bytes) -> Tuple[int, int]:
        """Read image pixel dimensions from image bytes."""
        from PIL import Image

        with Image.open(BytesIO(img_bytes)) as image:
            width_px, height_px = image.size

        if width_px <= 0 or height_px <= 0:
            raise ValueError("Image has invalid dimensions")

        return width_px, height_px

    def _get_previous_text_paragraph(self) -> Optional[Any]:
        """
        Return the immediate previous paragraph when it contains text.

        This is used for compact Mermaid diagrams that should stay visually
        attached to the preceding text block while remaining centered on
        their own line.
        """
        if not self.doc.paragraphs:
            return None

        previous_paragraph = self.doc.paragraphs[-1]
        if previous_paragraph.text.strip():
            return previous_paragraph

        return None

    def _calculate_mermaid_layout(
        self,
        img_bytes: bytes,
        mermaid_style: Dict[str, Any],
        can_follow_previous: bool = False,
    ) -> Dict[str, Any]:
        """
        Calculate a Mermaid image layout that fits the current page.

        The image keeps its aspect ratio and is constrained by both
        available page width and height. Very tall diagrams can be nudged
        onto a fresh page so Word has more room to place them cleanly.
        """
        img_width_px, img_height_px = self._get_image_pixel_size(img_bytes)
        aspect_ratio = img_height_px / img_width_px

        available_width = int(self._get_available_page_width())
        available_height = int(self._get_available_page_height())

        preferred_width_value = mermaid_style.get('width')
        preferred_width = (
            min(int(self._parse_length(preferred_width_value, default_unit='in')), available_width)
            if preferred_width_value
            else available_width
        )

        soft_max_height = int(
            available_height * float(mermaid_style.get('soft_max_height_ratio', 0.68))
        )
        hard_max_height = int(
            available_height * float(mermaid_style.get('hard_max_height_ratio', 0.82))
        )
        page_max_height = int(
            available_height * float(mermaid_style.get('page_max_height_ratio', 0.92))
        )
        page_break_threshold = int(
            available_height * float(mermaid_style.get('page_break_threshold_ratio', 0.9))
        )

        min_readable_width_value = mermaid_style.get('min_readable_width')
        min_readable_width = (
            int(self._parse_length(min_readable_width_value, default_unit='in'))
            if min_readable_width_value
            else None
        )

        preferred_height = int(round(preferred_width * aspect_ratio))
        oversized_strategy = str(mermaid_style.get('oversized_strategy', 'page')).lower()
        page_break_before = False
        page_layout_mode = False
        force_page_break_before = bool(
            mermaid_style.get('force_page_break_before_oversized', False)
        )

        target_width = preferred_width
        target_height = preferred_height

        if preferred_height > hard_max_height:
            target_height = hard_max_height
            target_width = int(round(target_height / aspect_ratio))

        if oversized_strategy == 'page' and (
            preferred_height > page_break_threshold
            or (min_readable_width is not None and target_width < min_readable_width)
        ):
            page_layout_mode = True
            page_break_before = force_page_break_before
            target_height = min(preferred_height, page_max_height)
            target_width = int(round(target_height / aspect_ratio))

        # Respect the configured preferred width and current page width.
        target_width = min(target_width, preferred_width, available_width)

        max_target_height = page_max_height if page_layout_mode else hard_max_height
        target_height = int(round(target_width * aspect_ratio))
        if target_height > max_target_height:
            target_height = max_target_height
            target_width = int(round(target_height / aspect_ratio))

        target_width = max(1, min(target_width, available_width))
        target_height = max(1, int(round(target_width * aspect_ratio)))

        if target_height > max_target_height:
            target_height = max_target_height
            target_width = max(1, int(round(target_height / aspect_ratio)))

        follow_previous_trigger_height = int(
            available_height * float(mermaid_style.get('follow_previous_trigger_height_ratio', 0.24))
        )
        follow_previous = False

        if can_follow_previous and not page_break_before and target_height <= follow_previous_trigger_height:
            compact_width_ratio = float(mermaid_style.get('follow_previous_width_ratio', 0.55))
            compact_width = max(1, int(round(available_width * compact_width_ratio)))

            if compact_width < target_width:
                target_width = compact_width
                target_height = max(1, int(round(target_width * aspect_ratio)))

            follow_previous = True

        return {
            'width': Emu(target_width),
            'height': Emu(target_height),
            'page_break_before': page_break_before,
            'soft_limit_exceeded': target_height > soft_max_height,
            'follow_previous_paragraph': follow_previous,
        }

    def _can_bind_mermaid_to_previous(self, paragraph: Any, mermaid_style: Dict[str, Any]) -> bool:
        """Return True when the previous text is short enough to bind to a Mermaid image."""
        max_chars = int(mermaid_style.get('keep_with_previous_max_chars', 80))
        if max_chars <= 0:
            return True

        return len(paragraph.text.strip()) <= max_chars

    def _add_block_image(
        self,
        img_bytes: bytes,
        width: Optional[Any] = None,
        height: Optional[Any] = None,
        alignment: str = 'center',
        space_before: Optional[Any] = None,
        space_after: Optional[Any] = None,
        keep_together: Optional[bool] = None,
        keep_with_next: Optional[bool] = None,
        page_break_before: Optional[bool] = None,
        widow_control: Optional[bool] = None,
    ) -> Any:
        """
        Add a block image as its own paragraph.

        Args:
            img_bytes: Image bytes to embed
            width: Optional target width
            height: Optional target height
            alignment: Paragraph alignment
            space_before: Optional spacing before
            space_after: Optional spacing after
            keep_together: Optional paragraph keep-together flag
            keep_with_next: Optional keep-with-next flag
            page_break_before: Optional page-break-before flag
            widow_control: Optional widow control flag
        """
        p = self.doc.add_paragraph()
        p.alignment = self._get_alignment(alignment)

        run = p.add_run()
        image_kwargs = {}
        if width is not None:
            image_kwargs['width'] = width
        if height is not None:
            image_kwargs['height'] = height
        shape = run.add_picture(BytesIO(img_bytes), **image_kwargs)

        if space_before is not None:
            p.paragraph_format.space_before = self._parse_length(space_before)
        if space_after is not None:
            p.paragraph_format.space_after = self._parse_length(space_after)

        if keep_together is not None:
            p.paragraph_format.keep_together = keep_together
        if keep_with_next is not None:
            p.paragraph_format.keep_with_next = keep_with_next
        if page_break_before is not None:
            p.paragraph_format.page_break_before = page_break_before
        if widow_control is not None:
            p.paragraph_format.widow_control = widow_control

        return p, shape

    def _resolve_image_path(self, url: str) -> Path:
        """Resolve an image URL against the Markdown base directory."""
        image_path = Path(unquote(url))
        if image_path.is_absolute():
            return image_path
        if self.base_dir is not None:
            return (self.base_dir / image_path).resolve()
        return image_path.resolve()

    def _embed_local_image(
        self,
        image_path: Path,
        display_path: Optional[str] = None,
        alignment: str = 'center',
    ) -> None:
        """Embed a local image file using the existing block-image helper."""
        visible_path = display_path or image_path.as_posix()
        if not image_path.exists():
            self.doc.add_paragraph(f'[Missing image: {visible_path}]')
            return

        suffix = image_path.suffix.lower().lstrip('.')
        supported_formats = {'png', 'jpg', 'jpeg', 'gif', 'bmp', 'tif', 'tiff'}
        if suffix and suffix not in supported_formats:
            self.doc.add_paragraph(f'[Unsupported image format: {visible_path}]')
            return

        try:
            img_bytes = image_path.read_bytes()
            width_px, height_px = self._get_image_pixel_size(img_bytes)
            available_width = int(self._get_available_page_width())
            available_height = int(self._get_available_page_height())
            dpi = 96
            try:
                from PIL import Image

                with Image.open(BytesIO(img_bytes)) as image:
                    dpi_info = image.info.get('dpi')
                    if dpi_info and dpi_info[0]:
                        dpi = float(dpi_info[0])
            except Exception:
                pass

            width_in = width_px / dpi
            height_in = height_px / dpi
            width_emu = int(Inches(width_in))
            height_emu = int(Inches(height_in))

            scale = min(
                1.0,
                available_width / width_emu if width_emu else 1.0,
                available_height / height_emu if height_emu else 1.0,
            )
            if scale < 1.0:
                width_emu = max(1, int(round(width_emu * scale)))
                height_emu = max(1, int(round(height_emu * scale)))

            self._add_block_image(
                img_bytes,
                width=Emu(width_emu),
                height=Emu(height_emu),
                alignment=alignment,
            )
        except Exception:
            self.doc.add_paragraph(f'[Failed to embed image: {visible_path}]')

    def image(self, token: Dict[str, Any], state: Any) -> str:
        """Render Markdown image tokens as embedded document images."""
        attrs = token.get('attrs') or {}
        url = attrs.get('url', '')
        if not url:
            return ''

        image_path = self._resolve_image_path(url)
        self._embed_local_image(image_path, display_path=url)
        return ''
    
    def _apply_paragraph_style(self, paragraph: Any, style: Dict[str, Any]) -> None:
        """Apply style to paragraph."""
        if 'font_name' in style:
            paragraph.style.font.name = style['font_name']
        
        if 'font_size' in style:
            paragraph.style.font.size = Pt(self._parse_font_size(style['font_size']))
        
        if 'font_color' in style:
            paragraph.style.font.color.rgb = self._parse_color(style['font_color'])
        
        if 'bold' in style and style['bold']:
            paragraph.style.font.bold = True
        
        if 'alignment' in style:
            paragraph.alignment = self._get_alignment(style['alignment'])
        
        # Spacing
        if 'space_before' in style:
            space_before = self._parse_font_size(style['space_before'])
            paragraph.paragraph_format.space_before = Pt(space_before)
        
        if 'space_after' in style:
            space_after = self._parse_font_size(style['space_after'])
            paragraph.paragraph_format.space_after = Pt(space_after)
    
    def render_children(self, token: Dict[str, Any], state: Any) -> list:
        """
        Render child tokens.
        
        Args:
            token: Token dictionary with 'children' key
            state: State object
            
        Returns:
            List of rendered strings
        """
        children = token.get('children', [])
        return [self.render_token(child, state) for child in children]
    
    def render_token(self, token: Dict[str, Any], state: Any) -> str:
        """
        Render a single token.
        
        Args:
            token: Token dictionary
            state: State object
            
        Returns:
            Rendered string
        """
        token_type = token['type']
        method_name = token_type
        
        # Get the render method for this token type
        method = getattr(self, method_name, None)
        
        if method:
            return method(token, state)
        
        # Fallback for unknown token types
        return ''
    
    def heading(self, token: Dict[str, Any], state: Any) -> str:
        """
        Render heading.
        
        Args:
            token: Token dictionary with 'children' and 'attrs' keys  
            state: BlockState object
            
        Returns:
            Empty string (content added to document)
        """
        # Extract level from token attributes
        level = token['attrs']['level']
        
        # Get text from children
        text = ''.join(self.render_children(token, state))
        
        # mistune uses levels 1-6, we support 1-4
        if level > 4:
            level = 4
        
        # Get style configuration
        style = self.styles.get_heading_style(level)
        
        # Add heading to document WITHOUT text first
        heading = self.doc.add_heading('', level=level)
        
        # Parse and add formatted text to heading
        self._add_formatted_text(heading, text, style)
        
        # Apply heading-level bold/italic from style configuration
        # These are SEPARATE from inline markdown formatting
        if 'bold' in style or 'italic' in style:
            for run in heading.runs:
                if 'bold' in style:
                    run.font.bold = style['bold']
                if 'italic' in style:
                    run.font.italic = style['italic']
        
        # Apply paragraph-level styles
        if 'alignment' in style:
            heading.alignment = self._get_alignment(style['alignment'])
        
        if 'space_before' in style:
            heading.paragraph_format.space_before = Pt(self._parse_font_size(style['space_before']))
        
        if 'space_after' in style:
            heading.paragraph_format.space_after = Pt(self._parse_font_size(style['space_after']))
        
        # Apply first line indent if specified
        if 'first_line_indent' in style and style['first_line_indent'] > 0:
            char_count = style['first_line_indent']
            font_size_pt = self._parse_font_size(style.get('font_size', '12pt'))
            heading.paragraph_format.first_line_indent = Pt(char_count * font_size_pt)
        
        return ''

    
    def paragraph(self, token: Dict[str, Any], state: Any) -> str:
        """
        Render paragraph.
        
        Args:
            token: Token dictionary with 'children' key
            state: BlockState object
            
        Returns:
            Empty string (content added to document)
        """
        # Get text from children
        text = ''.join(self.render_children(token, state))
        
        # Handle block math formulas (they appear as their own paragraphs)
        # Unicode brackets: 〔BLOCK_MATH_N〕
        text_stripped = text.strip()
        if text_stripped.startswith('〔BLOCK_MATH_') and text_stripped.endswith('〕'):
            idx_str = text_stripped.replace('〔BLOCK_MATH_', '').replace('〕', '')
            try:
                idx = int(idx_str)
                if idx < len(self.math_formulas.get('block', [])):
                    latex = self.math_formulas['block'][idx]
                    if self._add_block_math_formula(latex):
                        return ''
                    text = f'$$\\n{latex}\\n$$'
            except:
                pass
        
        # Skip empty paragraphs
        if not text or text.strip() == '':
            return ''

        if self._outline_mode_active():
            outline_level = self._match_outline_level(text)
            if outline_level is not None:
                self._render_outline_paragraph(text.strip(), outline_level)
                return ''
        
        # Get style configuration
        style = self.styles.get_paragraph_style()
        
        # Add paragraph to document
        p = self.doc.add_paragraph()
        
        # Parse inline formatting (bold, italic)
        self._add_formatted_text(p, text, style)
        self._apply_text_paragraph_format(p, style)
        
        return ''

    def _append_omml(self, paragraph: Any, latex: str, inline: bool) -> bool:
        """Append Word-native OMML math to a paragraph when conversion succeeds."""
        try:
            elements = self.omml_converter.latex_to_omml(latex, inline=inline)
        except Exception:
            return False

        if not elements:
            return False

        for element in elements:
            paragraph._p.append(element)
        return True

    def _apply_math_block_format(self, paragraph: Any) -> None:
        """Apply configured block math paragraph formatting."""
        math_style = self.styles.get_style('math_block', {})
        alignment = str(math_style.get('alignment', 'center')).lower()
        if alignment == 'left':
            paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
        elif alignment == 'right':
            paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
        else:
            paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        if 'space_before' in math_style:
            paragraph.paragraph_format.space_before = Pt(
                self._parse_font_size(math_style['space_before'])
            )
        else:
            paragraph.paragraph_format.space_before = Pt(6)

        if 'space_after' in math_style:
            paragraph.paragraph_format.space_after = Pt(
                self._parse_font_size(math_style['space_after'])
            )
        else:
            paragraph.paragraph_format.space_after = Pt(6)

    def _add_block_math_formula(self, latex: str) -> bool:
        """Add a display formula as OMML, falling back to the existing PNG path."""
        latex = latex.strip()
        if not latex:
            return False

        paragraph = self.doc.add_paragraph()
        self._apply_math_block_format(paragraph)
        if self._append_omml(paragraph, latex, inline=False):
            return True

        try:
            img_bytes = self.math_converter.latex_to_image(latex, inline=False)
            run = paragraph.add_run()
            math_style = self.styles.get_style('math_block', {})
            width_value = math_style.get('width')
            if width_value:
                width = self._parse_length(width_value)
                run.add_picture(BytesIO(img_bytes), width=width)
            else:
                run.add_picture(BytesIO(img_bytes), width=Inches(4))
            return True
        except Exception:
            paragraph._element.getparent().remove(paragraph._element)
            return False
    
    def _add_formatted_text(self, paragraph: Any, text: str, base_style: Dict[str, Any]) -> None:
        """
        Add text with inline formatting to paragraph.
        
        Args:
            paragraph: Paragraph object
            text: Text with HTML-like formatting tags
            base_style: Base style configuration
        """
        # mistune converts markdown to HTML tags in text
        # We need to parse <strong>, <em>, <code>, etc.
        
        # Replace HTML tags with markers
        text = text.replace('<strong>', '**START_BOLD**')
        text = text.replace('</strong>', '**END_BOLD**')
        text = text.replace('<em>', '**START_ITALIC**')
        text = text.replace('</em>', '**END_ITALIC**')
        text = text.replace('<code>', '**START_CODE**')
        text = text.replace('</code>', '**END_CODE**')
        
        # Split by markers and process (including math placeholders with Unicode brackets)
        parts = re.split(r'(\*\*START_BOLD\*\*|\*\*END_BOLD\*\*|\*\*START_ITALIC\*\*|\*\*END_ITALIC\*\*|\*\*START_CODE\*\*|\*\*END_CODE\*\*|〔INLINE_MATH_\d+〕|〔INLINE_MATH_DIRECT:.*?〕|〔BLOCK_MATH_\d+〕)', text)
        
        bold = False
        italic = False
        code = False
        
        # Get inline styles configuration
        inline_bold_style = self.styles.get_inline_style('bold')
        inline_italic_style = self.styles.get_inline_style('italic')
        inline_code_style = self.styles.get_inline_style('code')
        
        for part in parts:
            if part == '**START_BOLD**':
                bold = True
            elif part == '**END_BOLD**':
                bold = False
            elif part == '**START_ITALIC**':
                italic = True
            elif part == '**END_ITALIC**':
                italic = False
            elif part == '**START_CODE**':
                code = True
            elif part == '**END_CODE**':
                code = False
            elif part.startswith('〔INLINE_MATH_') and part.endswith('〕'):
                # Handle inline math formulas with Unicode brackets
                idx_str = part.replace('〔INLINE_MATH_', '').replace('〕', '')
                try:
                    idx = int(idx_str)
                    formulas = getattr(self, 'math_formulas', {'inline': []})
                    if idx < len(formulas.get('inline', [])):
                        latex = formulas['inline'][idx]
                        if not self._append_omml(paragraph, latex, inline=True):
                            try:
                                img_bytes = self.math_converter.latex_to_image(latex, inline=True)
                                run = paragraph.add_run()
                                run.add_picture(BytesIO(img_bytes), height=Inches(0.15))
                            except:
                                # Fallback: show LaTeX code
                                paragraph.add_run(f'${latex}$')
                except:
                    pass
            elif part.startswith('〔INLINE_MATH_DIRECT:') and part.endswith('〕'):
                latex = part.replace('〔INLINE_MATH_DIRECT:', '').replace('〕', '')
                if not self._append_omml(paragraph, latex, inline=True):
                    try:
                        img_bytes = self.math_converter.latex_to_image(latex, inline=True)
                        run = paragraph.add_run()
                        run.add_picture(BytesIO(img_bytes), height=Inches(0.15))
                    except:
                        paragraph.add_run(f'${latex}$')
            elif part.startswith('〔BLOCK_MATH_') and part.endswith('〕'):
                # Block math should not appear in inline text
                # This shouldn't happen with proper preprocessing
                pass
            elif part:  # Actual text content
                run = paragraph.add_run(part)
                
                # Apply code style if in code span
                if code:
                    # Apply code-specific styling
                    if inline_code_style.get('font_name'):
                        run.font.name = inline_code_style['font_name']
                        # Also set East Asian font
                        from docx.oxml import OxmlElement
                        from docx.oxml.ns import qn
                        run._element.rPr.rFonts.set(qn('w:eastAsia'), inline_code_style['font_name'])
                    
                    if inline_code_style.get('font_size'):
                        run.font.size = Pt(self._parse_font_size(inline_code_style['font_size']))
                    
                    if inline_code_style.get('font_color'):
                        run.font.color.rgb = self._parse_color(inline_code_style['font_color'])
                    
                    # Apply background color/shading
                    if inline_code_style.get('background'):
                        try:
                            from docx.oxml.ns import nsdecls
                            from docx.oxml import parse_xml
                            shading_elm = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(
                                nsdecls('w'), 
                                inline_code_style['background'].replace('#', '')
                            ))
                            run._element.get_or_add_rPr().append(shading_elm)
                        except:
                            pass  # Ignore if shading fails
                else:
                    # Apply base style - ALWAYS apply font settings
                    if 'font_name' in base_style:
                        run.font.name = base_style['font_name']
                        # Also set East Asian font for Chinese characters
                        from docx.oxml import OxmlElement
                        from docx.oxml.ns import qn
                        run._element.rPr.rFonts.set(qn('w:eastAsia'), base_style['font_name'])
                    
                    if 'font_size' in base_style:
                        run.font.size = Pt(self._parse_font_size(base_style['font_size']))
                    
                    if 'font_color' in base_style:
                        run.font.color.rgb = self._parse_color(base_style['font_color'])
                    
                    # Apply formatting on top of base style
                    if bold:
                        run.font.bold = True
                        # Apply bold color from inline config if specified
                        if inline_bold_style.get('font_color'):
                            run.font.color.rgb = self._parse_color(inline_bold_style['font_color'])
                    
                    if italic:
                        run.font.italic = True
                        # Apply italic color from inline config if specified
                        if inline_italic_style.get('font_color'):
                            run.font.color.rgb = self._parse_color(inline_italic_style['font_color'])
    
    def text(self, token: Dict[str, Any], state: Any) -> str:
        """Render plain text."""
        return token.get('raw', '')
    
    def strong(self, token: Dict[str, Any], state: Any) -> str:
        """Render strong (bold) text."""
        text = ''.join(self.render_children(token, state))
        return f"<strong>{text}</strong>"
    
    def emphasis(self, token: Dict[str, Any], state: Any) -> str:
        """Render emphasis (italic) text."""
        text = ''.join(self.render_children(token, state))
        return f"<em>{text}</em>"
    
    def codespan(self, token: Dict[str, Any], state: Any) -> str:
        """Render inline code span."""
        text = token.get('raw', '')
        return f"<code>{text}</code>"
    
    def inline_math(self, token: Dict[str, Any], state: Any) -> str:
        """
        Render inline LaTeX math formula.
        
        Args:
            token: Token with 'raw' containing LaTeX formula
            state: State object
            
        Returns:
            Marker string for later processing
        """
        latex = token.get('raw', '')
        if not latex:
            return ''

        return f'〔INLINE_MATH_DIRECT:{latex}〕'
    
    def block_math(self, token: Dict[str, Any], state: Any) -> str:
        """
        Render block LaTeX math formula.
        
        Args:
            token: Token with 'raw' containing LaTeX formula
            state: State object
            
        Returns:
            Empty string (content added to document)
        """
        latex = token.get('raw', '')
        if not latex:
            return ''

        if not self._add_block_math_formula(latex):
            p = self.doc.add_paragraph(f'$$\n{latex}\n$$')
            p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        return ''
    
    def block_code(self, token: Dict[str, Any], state: Any) -> str:
        """
        Render code block.
        
        Args:
            token: Token dictionary with 'raw' (code content) and optional 'attrs' (language)
            state: State object
            
        Returns:
            Empty string (content added to document)
        """
        # Get code content
        code_text = token.get('raw', '')
        attrs = token.get('attrs') or {}
        language_info = attrs.get('info', '') or ''
        language = language_info.split(None, 1)[0].lower() if language_info else ''
        
        if not code_text:
            return ''

        if language in {'latex', 'tex'}:
            if not self._add_block_math_formula(code_text):
                p = self.doc.add_paragraph(f'$$\n{code_text.strip()}\n$$')
                p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            return ''

        if language == 'mermaid':
            mermaid_style = self.styles.get_mermaid_style()
            try:
                img_bytes = self.mermaid_converter.mermaid_to_image(code_text)
                previous_paragraph = self._get_previous_text_paragraph()
                keep_with_previous = previous_paragraph is not None and bool(
                    mermaid_style.get('keep_with_previous', True)
                    and self._can_bind_mermaid_to_previous(previous_paragraph, mermaid_style)
                )
                layout = self._calculate_mermaid_layout(
                    img_bytes,
                    mermaid_style,
                    can_follow_previous=keep_with_previous,
                )

                if (
                    keep_with_previous
                    and previous_paragraph is not None
                    and layout.get('follow_previous_paragraph')
                ):
                    previous_paragraph.paragraph_format.keep_with_next = True

                self._add_block_image(
                    img_bytes,
                    width=layout['width'],
                    height=layout['height'],
                    alignment=mermaid_style.get('alignment', 'center'),
                    space_before=(
                        mermaid_style.get('follow_previous_space_before', '0pt')
                        if layout.get('follow_previous_paragraph') and keep_with_previous
                        else mermaid_style.get('space_before')
                    ),
                    space_after=mermaid_style.get('space_after'),
                    keep_together=mermaid_style.get('keep_together', True),
                    keep_with_next=mermaid_style.get('keep_with_next', False),
                    page_break_before=layout['page_break_before'],
                    widow_control=mermaid_style.get('widow_control', False),
                )
                return ''
            except Exception:
                # Fall back to the existing code block rendering if Mermaid render fails.
                pass
        
        # Get code block style
        code_style = self.styles.get_code_block_style()
        
        # Apply padding by adding spaces to each line
        if code_style.get('padding'):
            padding_pt = self._parse_font_size(code_style['padding'])
            # Approximate: 1 space ≈ 0.5em for monospace fonts
            # For 12pt padding with 10pt font: ~2-3 spaces
            font_size = self._parse_font_size(code_style.get('font_size', '10pt'))
            spaces_count = max(1, int(padding_pt / font_size * 2))
            padding_str = ' ' * spaces_count
            
            # Add padding to each line
            lines = code_text.split('\n')
            padded_lines = [padding_str + line for line in lines]
            code_text = '\n'.join(padded_lines)
        
        # Create paragraph for code block
        p = self.doc.add_paragraph()
        
        # Add code text
        run = p.add_run(code_text)
        
        # Apply font styling
        if code_style.get('font_name'):
            run.font.name = code_style['font_name']
            # Also set East Asian font
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
            run._element.rPr.rFonts.set(qn('w:eastAsia'), code_style['font_name'])
        
        if code_style.get('font_size'):
            run.font.size = Pt(self._parse_font_size(code_style['font_size']))
        
        # Apply bold and italic if specified
        if code_style.get('bold'):
            run.font.bold = True
        
        if code_style.get('italic'):
            run.font.italic = True
        
        if code_style.get('font_color'):
            run.font.color.rgb = self._parse_color(code_style['font_color'])
        
        # Apply background shading
        if code_style.get('background'):
            try:
                from docx.oxml.ns import nsdecls
                from docx.oxml import parse_xml
                shading_elm = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(
                    nsdecls('w'), 
                    code_style['background'].replace('#', '')
                ))
                p._element.get_or_add_pPr().append(shading_elm)
            except:
                pass  # Ignore if shading fails
        
        # Apply paragraph-level styling
        if code_style.get('line_spacing'):
            p.paragraph_format.line_spacing = code_style['line_spacing']
        
        if code_style.get('space_before'):
            p.paragraph_format.space_before = Pt(self._parse_font_size(code_style['space_before']))
        
        if code_style.get('space_after'):
            p.paragraph_format.space_after = Pt(self._parse_font_size(code_style['space_after']))
        
        return ''
    
    def linebreak(self, token: Dict[str, Any], state: Any) -> str:
        """Render line break."""
        return '\n'
    
    def softbreak(self, token: Dict[str, Any], state: Any) -> str:
        """Render soft break."""
        return ' '
    
    def blank_line(self, token: Dict[str, Any], state: Any) -> str:
        """Render blank line."""
        return ''
    
    def thematic_break(self, token: Dict[str, Any], state: Any) -> str:
        """Render horizontal rule."""
        document_style = self.styles.get_document_style()
        if document_style.get('ignore_thematic_breaks', True):
            return ''

        # Add a paragraph with a horizontal line
        p = self.doc.add_paragraph()
        p.add_run('_' * 50)
        return ''
    
    def list(self, token: Dict[str, Any], state: Any) -> str:
        """
        Render list (ordered or unordered).
        
        Args:
            token: Token dictionary with 'children' and 'attrs' keys
            state: BlockState object
            
        Returns:
            Empty string (content added to document)
        """
        # Get list type (ordered/unordered) and depth
        ordered = token['attrs'].get('ordered', False)
        depth = token.get('attrs', {}).get('depth', 0)
        list_num_id = self._create_ordered_list_num_id() if ordered else None
        
        # Process each list item
        children = token.get('children', [])
        for item_number, child in enumerate(children, start=1):
            if child['type'] == 'list_item':
                self._render_list_item(
                    child,
                    state,
                    ordered,
                    depth,
                    list_num_id,
                    item_number,
                )
        
        return ''
    
    def _render_list_item(
        self,
        token: Dict[str, Any],
        state: Any,
        ordered: bool,
        depth: int,
        list_num_id: Optional[int] = None,
        item_number: int = 1,
    ) -> None:
        """
        Render a single list item.
        
        Args:
            token: List item token
            state: State object
            ordered: Whether this is an ordered list
            depth: Nesting depth (0-based)
        """
        # Separate inline content from nested lists
        children = token.get('children', [])
        inline_children = []
        nested_lists = []
        
        for child in children:
            if child['type'] == 'list':
                nested_lists.append(child)
            else:
                inline_children.append(child)
        
        # Get text from inline children only. Loose Markdown lists wrap item
        # content in paragraph tokens; rendering those directly would create a
        # Normal paragraph before list formatting can be applied.
        text_parts = [
            self._render_list_item_text(child, state)
            for child in inline_children
        ]
        text = '\n'.join(part.strip() for part in text_parts if part and part.strip())
        
        # Only skip if there's no inline text AND no nested lists
        if (not text or text.strip() == '') and not nested_lists:
            return
        
        # If there's inline text, render it as a list item
        if text and text.strip():
            # Get list style
            list_style = self.styles.get_list_style()
            number_format = list_style.get('number_format', '1.')
            ordered_list_as_text = bool(list_style.get('ordered_list_as_text', False))
            
            # Add paragraph with appropriate style
            p = self.doc.add_paragraph()
            
            # Parse and add formatted text
            base_style = {
                'font_name': list_style.get('font_name', '宋体'),
                'font_size': list_style.get('font_size', '12pt')
            }
            if ordered and ordered_list_as_text:
                marker = self._format_ordered_list_marker(item_number, number_format)
                self._add_formatted_text(p, f'{marker} {text}', base_style)
            else:
                self._add_formatted_text(p, text, base_style)
            
            # Apply custom list formatting
            if ordered and not ordered_list_as_text:
                # Custom number format
                self._apply_numbered_list(p, number_format, depth, list_num_id)
            else:
                # Custom bullet character
                if not ordered:
                    bullet_char = list_style.get('bullet_char', '•')
                    self._apply_bulleted_list(p, bullet_char, depth)
            
            # Apply indentation for nested lists
            if depth > 0:
                from docx.shared import Inches
                indent_size = 0.5  # inches per level
                if 'indent_size' in list_style:
                    indent_str = list_style['indent_size']
                    if isinstance(indent_str, str) and indent_str.endswith('in'):
                        indent_size = float(indent_str[:-2])
                
                p.paragraph_format.left_indent = Inches(indent_size * (depth + 1))
            
            # Apply line spacing from list style
            if 'line_spacing' in list_style:
                p.paragraph_format.line_spacing = list_style['line_spacing']
            
            # Apply space after from list style
            if 'space_after' in list_style:
                p.paragraph_format.space_after = Pt(self._parse_font_size(list_style['space_after']))
        
        # Render nested lists with increased depth
        for nested_list in nested_lists:
            # Get list type from nested list
            nested_ordered = nested_list['attrs'].get('ordered', False)
            nested_depth = nested_list.get('attrs', {}).get('depth', depth + 1)
            nested_num_id = self._create_ordered_list_num_id() if nested_ordered else None
            # Process each child of the nested list
            for nested_item_number, child in enumerate(nested_list.get('children', []), start=1):
                if child['type'] == 'list_item':
                    self._render_list_item(
                        child,
                        state,
                        nested_ordered,
                        nested_depth,
                        nested_num_id,
                        nested_item_number,
                    )

    def _render_list_item_text(self, token: Dict[str, Any], state: Any) -> str:
        """
        Render list-item child content without writing a standalone paragraph.

        Mistune emits `block_text` for tight lists and `paragraph` for loose
        lists. Both should feed the list item paragraph instead of being
        rendered independently.
        """
        if token['type'] in {'paragraph', 'block_text'}:
            return ''.join(self.render_children(token, state))

        return self.render_token(token, state)

    def _format_ordered_list_marker(self, item_number: int, number_format: str) -> str:
        """
        Format a literal ordered-list marker using the configured pattern.

        Supported examples: `1.`, `1)`, `(1)`.
        """
        fmt = str(number_format or '1.').strip()
        if '1' in fmt:
            return fmt.replace('1', str(item_number), 1)
        return f'{item_number}.'
    
    def _apply_bulleted_list(self, paragraph, bullet_char: str, depth: int):
        """
        Apply custom bullet character to paragraph.
        
        Args:
            paragraph: Paragraph object
            bullet_char: Bullet character to use
            depth: Nesting depth
        """
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        
        # Get or create paragraph properties
        pPr = paragraph._element.get_or_add_pPr()
        
        # Create numbering properties
        numPr = OxmlElement('w:numPr')
        
        # Create indent level
        ilvl = OxmlElement('w:ilvl')
        ilvl.set(qn('w:val'), str(depth))
        numPr.append(ilvl)
        
        # Create numbering ID (use 1 for bullets)
        numId = OxmlElement('w:numId')
        numId.set(qn('w:val'), '1')
        numPr.append(numId)
        
        # Add to paragraph properties
        pPr.append(numPr)
        
        # Add custom bullet as text at the beginning (workaround)
        # This is a simpler approach than creating complex numbering definitions
        if bullet_char != '•':  # Only if not default
            # Remove the numbering and add bullet manually
            pPr.remove(numPr)
            
            # Insert bullet at the beginning of paragraph
            runs = paragraph.runs
            if runs:
                # Prepend bullet to first run
                first_run = runs[0]
                first_run.text = f"{bullet_char}\t{first_run.text}"
            else:
                # Add bullet run
                run = paragraph.add_run(f"{bullet_char}\t")
        else:
            # Use Word's default bullet style
            paragraph.style = 'List Bullet'
    
    def _apply_numbered_list(
        self,
        paragraph,
        number_format: str,
        depth: int,
        list_num_id: Optional[int],
    ):
        """
        Apply custom number format to paragraph.
        
        Args:
            paragraph: Paragraph object
            number_format: Number format string (e.g., "1.", "1)", "(1)")
            depth: Nesting depth
            list_num_id: Concrete numbering instance for this Markdown list block
        """
        paragraph.style = 'List Number'

        if list_num_id is not None:
            self._set_paragraph_numbering(paragraph, list_num_id, ilvl=0)

    
    def list_item(self, token: Dict[str, Any], state: Any) -> str:
        """
        Render list item.
        
        Note: This is called by list() method, but we need to provide it
        for the render_token dispatcher.
        
        Args:
            token: List item token
            state: State object
            
        Returns:
            Empty string
        """
        # When called directly, just render children and return text
        # The actual rendering is done by _render_list_item in list()
        return ''.join(self.render_children(token, state))
    
    def block_text(self, token: Dict[str, Any], state: Any) -> str:
        """
        Render block text (used in list items).
        
        Args:
            token: Block text token
            state: State object
            
        Returns:
            Rendered text string
        """
        # Block text is just a container, render its children
        return ''.join(self.render_children(token, state))
    
    def table(self, token: Dict[str, Any], state: Any) -> str:
        """
        Render table.
        
        Args:
            token: Table token with 'children' (table_head and table_body)
            state: State object
            
        Returns:
            Empty string (content added to document)
        """
        # Get table style
        table_style = self.styles.get_table_style()
        table_layout = self._resolve_table_layout(table_style)
        
        # Create table - we'll determine column count from header
        children = token.get('children', [])
        if not children:
            return ''
        
        # Find table_head to get column count
        col_count = 0
        table_head = None
        table_body = None
        
        for child in children:
            if child['type'] == 'table_head':
                table_head = child
                col_count = len(child.get('children', []))
            elif child['type'] == 'table_body':
                table_body = child
        
        if col_count == 0:
            return ''
        
        # Determine row count
        row_count = 1  # header row
        if table_body:
            row_count += len(table_body.get('children', []))
        
        # Create table
        table = self.doc.add_table(rows=row_count, cols=col_count)
        column_widths = self._calculate_table_column_widths(
            table_head,
            table_body,
            col_count,
            table_style,
        )
        self._set_table_column_widths(table, column_widths)
        
        self._apply_table_layout(table, table_style, table_layout)
        
        # Track current row index
        self._current_table = table
        self._current_row_idx = 0
        
        # Render table head
        if table_head:
            self.table_head(table_head, state)
        
        # Render table body
        if table_body:
            self.table_body(table_body, state)
        
        self._current_table = None
        self._current_row_idx = 0
        
        return ''
    
    def table_head(self, token: Dict[str, Any], state: Any) -> str:
        """Render table header."""
        cells = token.get('children', [])
        table_style = self.styles.get_table_style()
        table_layout = self._resolve_table_layout(table_style)
        header_alignment = str(table_style.get('header_alignment', 'center')).lower()
        header_vertical_alignment = str(table_style.get('header_vertical_alignment', 'center')).lower()
        
        if not self._current_table:
            return ''
        
        row = self._current_table.rows[self._current_row_idx]
        self._set_row_repeats_as_table_header(row)
        
        for col_idx, cell_token in enumerate(cells):
            if col_idx < len(row.cells):
                cell = row.cells[col_idx]
                text = ''.join(self.render_children(cell_token, state))
                
                # Clear existing paragraphs and add formatted text
                cell._element.clear_content()
                p = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
                
                # Parse and add formatted text
                base_style = {
                    'font_name': table_style.get('font_name', '宋体'),
                    'font_size': table_style.get('font_size', '11pt')
                }
                self._add_formatted_text(p, text, base_style)
                
                # Apply line spacing to table cell paragraph
                if 'line_spacing' in table_style:
                    p.paragraph_format.line_spacing = table_style['line_spacing']
                
                # Center header text by default, with an opt-in inherit mode.
                align = cell_token.get('attrs', {}).get('align') or table_style.get('alignment', 'left')
                if header_alignment == 'inherit':
                    p.alignment = self._get_alignment(str(align))
                else:
                    p.alignment = self._get_alignment(header_alignment)
                cell.vertical_alignment = self._get_vertical_alignment(header_vertical_alignment)
                self._apply_table_cell_margins(cell, table_style)
                
                # Apply header styling - make all runs bold
                for run in p.runs:
                    run.font.bold = True
                
                # Apply header background if specified
                if (
                    table_layout != TABLE_LAYOUT_THREE_LINE
                    and 'header_background' in table_style
                ):
                    try:
                        from docx.oxml.ns import nsdecls
                        from docx.oxml import parse_xml
                        shading_elm = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(
                            nsdecls('w'), 
                            table_style['header_background'].replace('#', '')
                        ))
                        cell._element.get_or_add_tcPr().append(shading_elm)
                    except:
                        pass  # Ignore if shading fails
        
        self._current_row_idx += 1
        return ''
    
    def table_body(self, token: Dict[str, Any], state: Any) -> str:
        """Render table body."""
        rows = token.get('children', [])
        
        for row_token in rows:
            if row_token['type'] == 'table_row':
                self.table_row(row_token, state)
        
        return ''
    
    def table_row(self, token: Dict[str, Any], state: Any) -> str:
        """Render table row."""
        cells = token.get('children', [])
        table_style = self.styles.get_table_style()
        table_layout = self._resolve_table_layout(table_style)
        
        if not self._current_table or self._current_row_idx >= len(self._current_table.rows):
            return ''
        
        row = self._current_table.rows[self._current_row_idx]
        
        # Check if alternating rows are enabled
        alternating_rows = (
            table_layout == TABLE_LAYOUT_ACCENT_GRID
            and table_style.get('alternating_rows', False)
        )
        
        for col_idx, cell_token in enumerate(cells):
            if col_idx < len(row.cells):
                cell = row.cells[col_idx]
                text = ''.join(self.render_children(cell_token, state))
                
                # Clear existing paragraphs and add formatted text
                cell._element.clear_content()
                p = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
                
                # Parse and add formatted text
                base_style = {
                    'font_name': table_style.get('font_name', '宋体'),
                    'font_size': table_style.get('font_size', '11pt')
                }
                self._add_formatted_text(p, text, base_style)

                if 'line_spacing' in table_style:
                    p.paragraph_format.line_spacing = table_style['line_spacing']
                
                # Apply column alignment, falling back to the configured table default.
                align = cell_token.get('attrs', {}).get('align') or table_style.get('alignment', 'left')
                p.alignment = self._get_alignment(str(align))

                vertical_alignment = str(table_style.get('vertical_alignment', 'center')).lower()
                cell.vertical_alignment = self._get_vertical_alignment(vertical_alignment)
                self._apply_table_cell_margins(cell, table_style)
                
                # Apply alternating row background (zebra striping)
                if alternating_rows:
                    # Row 0 is header, so data rows start from row 1
                    # Even rows: 2, 4, 6... Odd rows: 1, 3, 5...
                    is_even_row = (self._current_row_idx % 2) == 0
                    
                    if is_even_row and 'row_background_even' in table_style:
                        bg_color = table_style['row_background_even']
                    elif not is_even_row and 'row_background_odd' in table_style:
                        bg_color = table_style['row_background_odd']
                    else:
                        bg_color = None
                    
                    if bg_color:
                        try:
                            from docx.oxml.ns import nsdecls
                            from docx.oxml import parse_xml
                            shading_elm = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(
                                nsdecls('w'), 
                                bg_color.replace('#', '')
                            ))
                            cell._element.get_or_add_tcPr().append(shading_elm)
                        except:
                            pass  # Ignore if shading fails
        
        self._current_row_idx += 1
        return ''
    
    def table_cell(self, token: Dict[str, Any], state: Any) -> str:
        """Render table cell (called by render_children)."""
        return ''.join(self.render_children(token, state))
