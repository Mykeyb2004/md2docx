"""
Verify the latest md2docx history entry against the current converter.

This is a manual verification helper, not a regular automated test:
- reads ~/.md2docx/history.json
- picks the latest entry
- converts the source Markdown to a temporary DOCX
- reports Mermaid block counts, embedded images, and raw Mermaid leakage
"""
from __future__ import annotations

import argparse
import json
import re
import tempfile
import zipfile
from pathlib import Path

from docx import Document

from md2docx import Converter


MERMAID_BLOCK_RE = re.compile(
    r"(?P<fence>```|~~~)mermaid[^\n]*\n(?P<body>.*?)(?:\n(?P=fence))",
    re.DOTALL,
)


def load_history_entry(index: int = 0) -> dict:
    """Load a history entry from the user's persisted GUI history file."""
    history_path = Path.home() / ".md2docx" / "history.json"
    if not history_path.exists():
        raise FileNotFoundError(f"History file not found: {history_path}")

    items = json.loads(history_path.read_text(encoding="utf-8"))
    if not items:
        raise ValueError("History file is empty")
    if index >= len(items):
        raise IndexError(f"History index {index} out of range; total records: {len(items)}")
    return items[index]


def find_mermaid_blocks(markdown_text: str) -> list[str]:
    """Return Mermaid fenced block bodies from a Markdown document."""
    return [match.group("body").strip() for match in MERMAID_BLOCK_RE.finditer(markdown_text)]


def docx_media_summary(docx_path: Path) -> tuple[int, int]:
    """Return embedded media count and drawing count from a DOCX archive."""
    media_count = 0
    drawing_count = 0
    with zipfile.ZipFile(docx_path) as archive:
        media_count = len([name for name in archive.namelist() if name.startswith("word/media/")])
        document_xml = archive.read("word/document.xml").decode("utf-8")
        drawing_count = document_xml.count("<w:drawing>")
    return media_count, drawing_count


def leaked_mermaid_snippets(doc: Document, mermaid_blocks: list[str]) -> list[str]:
    """Detect Mermaid source that still appears as text in the rendered document."""
    doc_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    leaked = []
    for block in mermaid_blocks:
        lines = [line.strip() for line in block.splitlines() if line.strip()]
        if not lines:
            continue
        sample = lines[0]
        if sample and sample in doc_text:
            leaked.append(sample)
    return leaked


def verify_latest_history(index: int = 0) -> int:
    """Verify the selected history entry and print a readable summary."""
    entry = load_history_entry(index=index)
    input_path = Path(entry["input"])
    output_path = Path(entry["output"])

    print(f"History time: {entry['time']}")
    print(f"History status: {entry['status']}")
    print(f"Input exists: {input_path.exists()} -> {input_path}")
    print(f"History output exists now: {output_path.exists()} -> {output_path}")

    if not input_path.exists():
        print("Result: input file is missing, cannot verify conversion.")
        return 2

    markdown_text = input_path.read_text(encoding="utf-8")
    mermaid_blocks = find_mermaid_blocks(markdown_text)
    print(f"Mermaid block count in input: {len(mermaid_blocks)}")

    with tempfile.TemporaryDirectory(prefix="md2docx-verify-") as tmp_dir:
        temp_docx = Path(tmp_dir) / "verify-output.docx"
        Converter().convert(str(input_path), str(temp_docx))

        doc = Document(str(temp_docx))
        media_count, drawing_count = docx_media_summary(temp_docx)
        leaked = leaked_mermaid_snippets(doc, mermaid_blocks)

        print(f"Temporary output: {temp_docx}")
        print(f"Paragraph count: {len(doc.paragraphs)}")
        print(f"Embedded media files: {media_count}")
        print(f"Drawing elements: {drawing_count}")
        print(f"Leaked Mermaid snippets: {len(leaked)}")

        if leaked:
            print("Leaked samples:")
            for sample in leaked[:5]:
                print(f"  - {sample}")

        if mermaid_blocks and drawing_count >= len(mermaid_blocks) and not leaked:
            print("Result: Mermaid rendering looks successful.")
            return 0

        if not mermaid_blocks:
            print("Result: no Mermaid blocks in the latest history entry; conversion itself succeeded.")
            return 0

        print("Result: conversion completed, but Mermaid verification did not fully pass.")
        return 1


def main() -> int:
    """CLI entry point."""
    parser = argparse.ArgumentParser(description="Verify the latest md2docx history entry.")
    parser.add_argument(
        "--index",
        type=int,
        default=0,
        help="History index to verify; 0 means the latest record.",
    )
    args = parser.parse_args()
    return verify_latest_history(index=args.index)


if __name__ == "__main__":
    raise SystemExit(main())
