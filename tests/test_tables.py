"""
Test table rendering functionality.
"""
import pytest
from pathlib import Path
import tempfile
from md2docx import Converter
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn


def _table_grid_widths(table):
    """Return table grid column widths in twips."""
    grid_cols = table._tbl.tblGrid.gridCol_lst
    widths = [col.get(qn("w:w")) for col in grid_cols]
    assert all(width is not None for width in widths)
    return [int(width) for width in widths]


def _width_spread(widths):
    """Return the absolute spread between the widest and narrowest columns."""
    return max(widths) - min(widths)


def _cell_margin_twips(cell, side):
    """Return one explicit table cell margin in twips."""
    margins = cell._tc.tcPr.first_child_found_in("w:tcMar")
    assert margins is not None
    margin = margins.find(qn(f"w:{side}"))
    assert margin is not None
    value = margin.get(qn("w:w"))
    assert value is not None
    return int(value)


def _row_repeats_as_header(row):
    """Return True when a Word table row is marked to repeat as a header."""
    tr_pr = row._tr.trPr
    if tr_pr is None:
        return False
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        return False
    return tbl_header.get(qn("w:val"), "true") != "false"


def test_simple_table():
    """Test simple table conversion."""
    converter = Converter()
    
    md_content = """# Table Test

| Column 1 | Column 2 | Column 3 |
|----------|----------|----------|
| A1       | B1       | C1       |
| A2       | B2       | C2       |
"""
    
    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
        output_path = f.name
    
    try:
        converter.convert_string(md_content, output_path)
        assert Path(output_path).exists()
        
        doc = Document(output_path)
        # Should have 1 table
        assert len(doc.tables) >= 1
        
        # Check table dimensions
        table = doc.tables[0]
        assert len(table.rows) == 3  # header + 2 data rows
        assert len(table.columns) == 3
    finally:
        Path(output_path).unlink(missing_ok=True)


def test_table_with_chinese():
    """Test table with Chinese content."""
    converter = Converter()
    
    md_content = """
| 姓名 | 年龄 | 城市 |
|------|------|------|
| 张三 | 25   | 北京 |
| 李四 | 30   | 上海 |
"""
    
    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
        output_path = f.name
    
    try:
        converter.convert_string(md_content, output_path)
        assert Path(output_path).exists()
        
        doc = Document(output_path)
        assert len(doc.tables) >= 1
        
        table = doc.tables[0]
        # Check header
        assert "姓名" in table.rows[0].cells[0].text
    finally:
        Path(output_path).unlink(missing_ok=True)


def test_default_table_column_widths_are_content_weighted():
    """Default table widths should give longer text columns more room."""
    converter = Converter()

    md_content = """
| 阶段 | 核心任务 | 成果 |
|------|----------|------|
| 采集 | 跨源提取、字段标准化、重复记录识别与基础数据质量校验 | 数据集 |
| 审核 | 表内表间逻辑比对、异常规则核查、疑似问题回溯定位 | 报告 |
"""

    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
        output_path = f.name

    try:
        converter.convert_string(md_content, output_path)
        doc = Document(output_path)
        widths = _table_grid_widths(doc.tables[0])

        assert len(widths) == 3
        assert widths[1] > widths[0]
        assert widths[1] > widths[2]
        assert len(set(widths)) > 1
    finally:
        Path(output_path).unlink(missing_ok=True)


def test_balanced_table_column_widths_are_more_even_than_default():
    """Balanced strategy should reduce the width spread for the same table."""
    md_content = """
| 阶段 | 核心任务 | 成果 |
|------|----------|------|
| 采集 | 跨源提取、字段标准化、重复记录识别与基础数据质量校验 | 数据集 |
| 审核 | 表内表间逻辑比对、异常规则核查、疑似问题回溯定位 | 报告 |
"""

    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as default_file:
        default_output_path = default_file.name
    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as balanced_file:
        balanced_output_path = balanced_file.name

    try:
        Converter().convert_string(md_content, default_output_path)
        Converter(
            config_override={"table": {"column_width_strategy": "balanced"}}
        ).convert_string(md_content, balanced_output_path)

        default_doc = Document(default_output_path)
        balanced_doc = Document(balanced_output_path)
        default_widths = _table_grid_widths(default_doc.tables[0])
        balanced_widths = _table_grid_widths(balanced_doc.tables[0])

        assert len(balanced_widths) == 3
        assert balanced_widths[1] > balanced_widths[0]
        assert balanced_widths[1] > balanced_widths[2]
        assert _width_spread(balanced_widths) < _width_spread(default_widths)
    finally:
        Path(default_output_path).unlink(missing_ok=True)
        Path(balanced_output_path).unlink(missing_ok=True)


def test_mixed_content_with_table():
    """Test document with mixed content."""
    converter = Converter()
    
    md_content = """# Report

## Summary

This is a summary paragraph.

## Data Table

| Metric | Value | Change |
|--------|-------|--------|
| Users  | 1000  | +10%   |
| Revenue| $5000 | +15%   |

## Conclusion

End of report.
"""
    
    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
        output_path = f.name
    
    try:
        converter.convert_string(md_content, output_path)
        assert Path(output_path).exists()
        
        doc = Document(output_path)
        assert len(doc.tables) >= 1
        assert len(doc.paragraphs) > 3
    finally:
        Path(output_path).unlink(missing_ok=True)


def test_table_header_is_centered_both_horizontally_and_vertically():
    """Test table headers default to centered alignment in both directions."""
    converter = Converter()

    md_content = """
| 类型 | 主要特征 |
|------|----------|
| A | B |
"""

    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
        output_path = f.name

    try:
        converter.convert_string(md_content, output_path)
        doc = Document(output_path)
        table = doc.tables[0]
        header_cell = table.rows[0].cells[0]
        header_paragraph = header_cell.paragraphs[0]

        assert header_paragraph.alignment == WD_PARAGRAPH_ALIGNMENT.CENTER
        assert header_cell.vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.CENTER
    finally:
        Path(output_path).unlink(missing_ok=True)


def test_table_header_row_repeats_on_each_page():
    """Test table header rows are marked for Word header repetition."""
    converter = Converter()

    md_content = """
| 类型 | 主要特征 |
|------|----------|
| A | B |
"""

    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
        output_path = f.name

    try:
        converter.convert_string(md_content, output_path)
        doc = Document(output_path)
        table = doc.tables[0]

        assert _row_repeats_as_header(table.rows[0])
        assert not _row_repeats_as_header(table.rows[1])
    finally:
        Path(output_path).unlink(missing_ok=True)


def test_table_body_cells_are_vertically_centered_with_line_spacing():
    """Test table body cells default to vertical center while keeping line spacing."""
    converter = Converter()

    md_content = """
| 类型 | 主要特征 |
|------|----------|
| 居民 | 反映居民识灾避险、预警响应和家庭应急准备。 |
"""

    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
        output_path = f.name

    try:
        converter.convert_string(md_content, output_path)
        doc = Document(output_path)
        table = doc.tables[0]
        body_cell = table.rows[1].cells[0]
        body_paragraph = body_cell.paragraphs[0]

        assert body_cell.vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.CENTER
        assert body_paragraph.paragraph_format.line_spacing == 1.5
        assert _cell_margin_twips(body_cell, "top") == 60
        assert _cell_margin_twips(body_cell, "bottom") == 60
    finally:
        Path(output_path).unlink(missing_ok=True)


def test_missing_separator_table_is_not_auto_fixed_by_default():
    """Malformed pipe blocks should stay plain text unless auto-fix is enabled."""
    converter = Converter()

    md_content = """
表格标题
| 列1 | 列2 | 列3 |
| A | B | C |
| D | E | F |
"""

    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
        output_path = f.name

    try:
        converter.convert_string(md_content, output_path)
        doc = Document(output_path)

        assert len(doc.tables) == 0
        assert any('|' in p.text for p in doc.paragraphs)
    finally:
        Path(output_path).unlink(missing_ok=True)


def test_missing_separator_table_can_be_auto_fixed_when_enabled():
    """Auto-fix should insert a separator row for high-confidence malformed tables."""
    converter = Converter(config_override={"document": {"auto_fix_tables": True}})

    md_content = """
表格标题
| 列1 | 列2 | 列3 |
| A | B | C |
| D | E | F |
"""

    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
        output_path = f.name

    try:
        converter.convert_string(md_content, output_path)
        doc = Document(output_path)

        assert len(doc.tables) == 1
        table = doc.tables[0]
        assert len(table.rows) == 3
        assert table.rows[0].cells[0].text == "列1"
        assert table.rows[1].cells[1].text == "B"
        assert not any('| 列1 | 列2 | 列3 |' in p.text for p in doc.paragraphs)
    finally:
        Path(output_path).unlink(missing_ok=True)


def test_captioned_malformed_table_can_be_split_and_fixed():
    """Single-cell caption lines should be preserved as text before the repaired table."""
    converter = Converter(config_override={"document": {"auto_fix_tables": True}})

    md_content = """
| 名录库维护关键节点与交付对照表 |
| 维护阶段 | 核心任务 | 输出成果 |
| 数据采集与清洗 | 跨源提取 | 清洗后基础数据集 |
| 数据审核与逻辑校验 | 表内表间逻辑比对 | 审核报告与异常工单 |
"""

    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
        output_path = f.name

    try:
        converter.convert_string(md_content, output_path)
        doc = Document(output_path)

        assert len(doc.tables) == 1
        assert any(p.text == "名录库维护关键节点与交付对照表" for p in doc.paragraphs)
        table = doc.tables[0]
        assert table.rows[0].cells[0].text == "维护阶段"
        assert table.rows[1].cells[2].text == "清洗后基础数据集"
    finally:
        Path(output_path).unlink(missing_ok=True)
