"""
Tests for code block and inline code rendering.
"""
from pathlib import Path
import pytest
from docx import Document

from md2docx import Converter


def test_inline_code():
    """Test inline code rendering."""
    md_content = """
这是一段包含 `inline code` 的文本。

可以在段落中使用 `变量名` 或者 `function()` 这样的代码。
"""
    
    converter = Converter()
    doc = converter.to_document(md_content)
    
    # Check that document has paragraphs
    assert len(doc.paragraphs) > 0
    
    # Check that code has been rendered
    # Note: We can't easily verify the exact styling, but we can check the text is there
    full_text = '\n'.join([p.text for p in doc.paragraphs])
    assert 'inline code' in full_text
    assert '变量名' in full_text


def test_block_code_simple():
    """Test simple code block rendering."""
    md_content = """
这是一段文本。

```
def hello():
    print("Hello, World!")
```

这是代码块后的文本。
"""
    
    converter = Converter()
    doc = converter.to_document(md_content)
    
    # Check that document has paragraphs
    assert len(doc.paragraphs) >= 3
    
    # Check that code text is present
    full_text = '\n'.join([p.text for p in doc.paragraphs])
    assert 'def hello():' in full_text
    assert 'print("Hello, World!")' in full_text


def test_block_code_with_language():
    """Test code block with language marker."""
    md_content = """
Python 代码示例：

```python
def fibonacci(n):
    if n <= 1:
        return n
    return fibonacci(n-1) + fibonacci(n-2)
```

JavaScript 代码示例：

```javascript
function factorial(n) {
    return n <= 1 ? 1 : n * factorial(n - 1);
}
```
"""
    
    converter = Converter()
    doc = converter.to_document(md_content)
    
    # Check that both code blocks are rendered
    full_text = '\n'.join([p.text for p in doc.paragraphs])
    assert 'def fibonacci' in full_text
    assert 'function factorial' in full_text


def test_mixed_code_and_text():
    """Test document with mixed code blocks and inline code."""
    md_content = """
# 代码示例文档

## 行内代码

可以使用 `print()` 函数输出内容，或者使用 `len()` 获取长度。

## 代码块

完整的函数示例：

```python
def greet(name):
    message = f"Hello, {name}!"
    print(message)
    return message
```

调用方式：`greet("World")`
"""
    
    converter = Converter()
    doc = converter.to_document(md_content)
    
    # Check headings
    assert any('代码示例文档' in p.text for p in doc.paragraphs)
    
    # Check inline code
    full_text = '\n'.join([p.text for p in doc.paragraphs])
    assert 'print()' in full_text
    assert 'len()' in full_text
    
    # Check code block
    assert 'def greet(name):' in full_text


def test_code_block_with_chinese():
    """Test code block containing Chinese comments."""
    md_content = """
```python
# 这是中文注释
def add(a, b):
    # 返回两个数的和
    return a + b

result = add(1, 2)  # 结果是3
```
"""
    
    converter = Converter()
    doc = converter.to_document(md_content)
    
    full_text = '\n'.join([p.text for p in doc.paragraphs])
    assert '这是中文注释' in full_text
    assert '返回两个数的和' in full_text


def test_multiple_code_blocks():
    """Test document with multiple code blocks."""
    md_content = """
第一个代码块：

```
code block 1
```

第二个代码块：

```
code block 2
```

第三个代码块：

```
code block 3
```
"""
    
    converter = Converter()
    doc = converter.to_document(md_content)
    
    full_text = '\n'.join([p.text for p in doc.paragraphs])
    assert 'code block 1' in full_text
    assert 'code block 2' in full_text
    assert 'code block 3' in full_text


def test_code_in_list():
    """Test inline code within list items."""
    md_content = """
常用函数：

* `print()` - 输出函数
* `len()` - 长度函数
* `range()` - 范围函数
"""
    
    converter = Converter()
    doc = converter.to_document(md_content)
    
    full_text = '\n'.join([p.text for p in doc.paragraphs])
    assert 'print()' in full_text
    assert 'len()' in full_text
    assert 'range()' in full_text


def test_empty_code_block():
    """Test empty code block."""
    md_content = """
空代码块：

```
```

后续文本。
"""
    
    converter = Converter()
    doc = converter.to_document(md_content)
    
    # Should not crash, just skip empty code block
    assert len(doc.paragraphs) >= 2
