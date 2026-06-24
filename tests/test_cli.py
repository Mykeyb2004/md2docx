"""
Tests for the md2docx command-line interface.
"""
import sys
from pathlib import Path

from docx import Document

from md2docx.cli import main


def test_output_dir_saves_docx_with_input_stem(tmp_path, monkeypatch):
    """--output-dir should save the converted file under the requested directory."""
    input_path = tmp_path / "report.md"
    output_dir = tmp_path / "converted"
    input_path.write_text("# Report\n\nCLI output directory test.", encoding="utf-8")

    monkeypatch.setattr(
        sys,
        "argv",
        ["md2docx", str(input_path), "--output-dir", str(output_dir)],
    )

    main()

    output_path = output_dir / "report.docx"
    assert output_path.exists()
    doc = Document(output_path)
    assert doc.paragraphs[0].text == "Report"


def test_output_file_saves_docx_to_requested_path(tmp_path, monkeypatch):
    """--output-file should save the converted document to the requested path."""
    input_path = tmp_path / "report.md"
    output_path = tmp_path / "custom.docx"
    input_path.write_text("# Report\n\nCLI output file test.", encoding="utf-8")

    monkeypatch.setattr(
        sys,
        "argv",
        ["md2docx", str(input_path), "--output-file", str(output_path)],
    )

    main()

    assert output_path.exists()
    doc = Document(output_path)
    assert doc.paragraphs[0].text == "Report"


def test_existing_output_file_requires_overwrite(tmp_path, monkeypatch, capsys):
    """CLI should protect an existing output file unless --overwrite is set."""
    input_path = tmp_path / "report.md"
    output_path = tmp_path / "custom.docx"
    input_path.write_text("# Report\n\nNew content.", encoding="utf-8")
    output_path.write_text("keep me", encoding="utf-8")

    monkeypatch.setattr(
        sys,
        "argv",
        ["md2docx", str(input_path), "--output-file", str(output_path)],
    )

    try:
        main()
    except SystemExit as exc:
        assert exc.code == 2
    else:
        raise AssertionError("Expected SystemExit when output exists")

    captured = capsys.readouterr()
    assert "output already exists" in captured.err
    assert "--overwrite" in captured.err
    assert output_path.read_text(encoding="utf-8") == "keep me"


def test_overwrite_allows_existing_output_file(tmp_path, monkeypatch):
    """--overwrite should allow replacing an existing output file."""
    input_path = tmp_path / "report.md"
    output_path = tmp_path / "custom.docx"
    input_path.write_text("# Report\n\nReplacement content.", encoding="utf-8")
    output_path.write_text("replace me", encoding="utf-8")

    monkeypatch.setattr(
        sys,
        "argv",
        [
            "md2docx",
            str(input_path),
            "--output-file",
            str(output_path),
            "--overwrite",
        ],
    )

    main()

    doc = Document(output_path)
    assert doc.paragraphs[0].text == "Report"


def test_legacy_output_alias_still_saves_docx(tmp_path, monkeypatch):
    """--output should remain accepted for existing scripts."""
    input_path = tmp_path / "report.md"
    output_path = tmp_path / "legacy.docx"
    input_path.write_text("# Report\n\nLegacy CLI output alias test.", encoding="utf-8")

    monkeypatch.setattr(
        sys,
        "argv",
        ["md2docx", str(input_path), "--output", str(output_path)],
    )

    main()

    assert output_path.exists()
    doc = Document(output_path)
    assert doc.paragraphs[0].text == "Report"


def test_output_file_and_output_dir_are_mutually_exclusive(tmp_path, monkeypatch, capsys):
    """CLI should reject ambiguous output path options."""
    input_path = tmp_path / "report.md"
    input_path.write_text("# Report\n", encoding="utf-8")

    monkeypatch.setattr(
        sys,
        "argv",
        [
            "md2docx",
            str(input_path),
            "--output-file",
            str(tmp_path / "custom.docx"),
            "--output-dir",
            str(tmp_path / "converted"),
        ],
    )

    try:
        main()
    except SystemExit as exc:
        assert exc.code == 2
    else:
        raise AssertionError("Expected SystemExit when output options conflict")

    captured = capsys.readouterr()
    assert "usage: md2docx" in captured.err
    assert "choose one: --output-file FILE or --output-dir DIR" in captured.err


def test_template_and_style_file_are_mutually_exclusive(tmp_path, monkeypatch, capsys):
    """CLI should reject ambiguous style source options."""
    input_path = tmp_path / "report.md"
    style_path = tmp_path / "style.yaml"
    input_path.write_text("# Report\n", encoding="utf-8")
    style_path.write_text("paragraph:\n  font_name: TestFont\n", encoding="utf-8")

    monkeypatch.setattr(
        sys,
        "argv",
        [
            "md2docx",
            str(input_path),
            "--template",
            "default",
            "--style-file",
            str(style_path),
        ],
    )

    try:
        main()
    except SystemExit as exc:
        assert exc.code == 2
    else:
        raise AssertionError("Expected SystemExit when style options conflict")

    captured = capsys.readouterr()
    assert "usage: md2docx" in captured.err
    assert "choose one: --template NAME or --style-file FILE" in captured.err


def test_help_prefers_output_file_name(monkeypatch, capsys):
    """Help should present --output-file as the primary output file option."""
    monkeypatch.setattr(sys, "argv", ["md2docx", "--help"])

    try:
        main()
    except SystemExit as exc:
        assert exc.code == 0
    else:
        raise AssertionError("Expected SystemExit for --help")

    captured = capsys.readouterr()
    assert "-o FILE, --output-file FILE" in captured.out
    assert "--output OUTPUT" not in captured.out
    assert "--overwrite" in captured.out
    assert "-s FILE, --style-file FILE, --style-config FILE" in captured.out


def test_directory_input_recurses_and_flattens_output_names(tmp_path, monkeypatch):
    """Directory input should convert all nested Markdown files into one output folder."""
    source_root = tmp_path / "source"
    output_dir = tmp_path / "converted"
    (source_root / "a").mkdir(parents=True)
    (source_root / "b").mkdir(parents=True)
    (source_root / "a" / "report.md").write_text("# A\n\nFirst.", encoding="utf-8")
    (source_root / "b" / "report.md").write_text("# B\n\nSecond.", encoding="utf-8")
    (source_root / "ignore.txt").write_text("ignore me", encoding="utf-8")

    monkeypatch.setattr(
        sys,
        "argv",
        ["md2docx", str(source_root), "--output-dir", str(output_dir)],
    )

    main()

    first_output = output_dir / "a_report.docx"
    second_output = output_dir / "b_report.docx"
    assert first_output.exists()
    assert second_output.exists()
    assert not (output_dir / "ignore.docx").exists()
    assert Document(first_output).paragraphs[0].text == "A"
    assert Document(second_output).paragraphs[0].text == "B"


def test_directory_output_refuses_existing_target_without_overwrite(tmp_path, monkeypatch, capsys):
    """Directory conversion should not partially overwrite existing target files."""
    source_root = tmp_path / "source"
    output_dir = tmp_path / "converted"
    (source_root / "a").mkdir(parents=True)
    output_dir.mkdir()
    (source_root / "a" / "report.md").write_text("# A\n\nFirst.", encoding="utf-8")
    existing_output = output_dir / "a_report.docx"
    existing_output.write_text("keep me", encoding="utf-8")

    monkeypatch.setattr(
        sys,
        "argv",
        ["md2docx", str(source_root), "--output-dir", str(output_dir)],
    )

    try:
        main()
    except SystemExit as exc:
        assert exc.code == 2
    else:
        raise AssertionError("Expected SystemExit when directory output exists")

    captured = capsys.readouterr()
    assert "output already exists" in captured.err
    assert "a_report.docx" in captured.err
    assert existing_output.read_text(encoding="utf-8") == "keep me"
