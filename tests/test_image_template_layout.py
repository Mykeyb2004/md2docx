"""Images must reserve their full height even in fixed-line-height templates."""
from io import BytesIO

import pytest
from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Pt
from PIL import Image

from md2docx import Converter
from md2docx.mermaid_converter import MermaidConverter


@pytest.mark.parametrize('image_kind', ['mermaid', 'markdown'])
def test_block_image_overrides_fixed_template_line_height(tmp_path, monkeypatch, image_kind):
    template = Document()
    template.styles['Normal'].paragraph_format.line_spacing = Pt(28)
    template_path = tmp_path / 'fixed-spacing-template.docx'
    template.save(template_path)

    png = BytesIO()
    Image.new('RGB', (600, 700), 'blue').save(png, format='PNG')
    image_path = tmp_path / 'diagram.png'
    image_path.write_bytes(png.getvalue())
    if image_kind == 'mermaid':
        # Replace only the external diagram renderer; exercise real Word insertion.
        monkeypatch.setattr(MermaidConverter, 'mermaid_to_image', lambda self, code: png.getvalue())
        image_markdown = '```mermaid\ngraph TD\nA --> B\n```'
    else:
        image_markdown = '![Diagram](diagram.png)'

    output = tmp_path / 'output.docx'
    Converter(word_template=str(template_path)).convert_string(
        f'Before the diagram.\n\n{image_markdown}\n\nAfter the diagram.',
        str(output),
        base_dir=tmp_path,
    )
    doc = Document(output)
    assert len(doc.inline_shapes) == 1
    assert doc.inline_shapes[0].height > Pt(28)
    image_paragraph = next(p for p in doc.paragraphs if p._p.xpath('.//w:drawing'))
    line_rule = image_paragraph.paragraph_format.line_spacing_rule
    if line_rule is None:
        line_rule = image_paragraph.style.paragraph_format.line_spacing_rule
    assert line_rule != WD_LINE_SPACING.EXACTLY, 'Fixed template spacing clips the image'

    assert doc.styles['Normal'].paragraph_format.line_spacing == Pt(28)
    assert doc.paragraphs[0].text == 'Before the diagram.'
    assert doc.paragraphs[-1].text == 'After the diagram.'
