"""
Test the full conversion workflow.
"""
import hashlib
import pytest
import xml.etree.ElementTree as ET
from pathlib import Path
import tempfile
from zipfile import ZipFile
from md2docx import Converter
from docx.document import Document as DocxDocument
from docx import Document
from docx.shared import Inches, RGBColor
from PIL import Image


def _write_sample_png(path: Path) -> None:
    """Create a tiny PNG image for document embedding tests."""
    path.parent.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (32, 24), (40, 120, 200))
    image.save(path, format="PNG")


def _docx_media_files(path: Path) -> list[str]:
    with ZipFile(path) as archive:
        return [
            name
            for name in archive.namelist()
            if name.startswith("word/media/")
        ]


def _docx_xml(path: Path, member: str) -> ET.Element:
    with ZipFile(path) as archive:
        return ET.fromstring(archive.read(member))


def _word_template_parts(path: Path) -> dict[str, str]:
    """Hash template-owned header/footer XML, relationships, and media parts."""
    prefixes = ("word/header", "word/footer", "word/media/")
    relationship_names = {"word/_rels/document.xml.rels"}
    with ZipFile(path) as archive:
        relationship_names.update(
            name
            for name in archive.namelist()
            if name.startswith("word/_rels/header")
            or name.startswith("word/_rels/footer")
        )
        names = {
            name
            for name in archive.namelist()
            if name.startswith(prefixes) or name in relationship_names
        }
        return {
            name: hashlib.sha256(archive.read(name)).hexdigest()
            for name in sorted(names)
        }


def _create_word_template(path: Path, image_path: Path, *, sections: int = 1) -> None:
    """Create a DOCX fixture with formatted primary, first-page, and even headers/footers."""
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.95)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.3)
    section.different_first_page_header_footer = True
    document.settings.odd_and_even_pages_header_footer = True

    primary_header = section.header.paragraphs[0]
    primary_header.alignment = 2
    primary_run = primary_header.add_run("Primary header")
    primary_run.bold = True
    primary_run.font.color.rgb = RGBColor(0x12, 0x34, 0x56)
    primary_header.add_run().add_picture(str(image_path), width=Inches(0.25))
    section.first_page_header.paragraphs[0].add_run("First-page header").italic = True
    section.even_page_header.paragraphs[0].add_run("Even-page header").underline = True

    section.footer.paragraphs[0].add_run("Primary footer").bold = True
    section.first_page_footer.paragraphs[0].add_run("First-page footer")
    section.even_page_footer.paragraphs[0].add_run("Even-page footer")
    document.add_paragraph("Template body placeholder")
    for _ in range(sections - 1):
        document.add_section()
    document.save(path)


def _remove_numbering_part(path: Path) -> None:
    """Strip numbering XML and its document relationship from a DOCX fixture."""
    temporary_path = path.with_suffix(".without-numbering.docx")
    with ZipFile(path) as source, ZipFile(temporary_path, "w") as target:
        for item in source.infolist():
            if item.filename in {"word/numbering.xml", "word/_rels/document.xml.rels"}:
                if item.filename == "word/_rels/document.xml.rels":
                    relationships = source.read(item.filename).decode("utf-8")
                    relationships = relationships.replace(
                        '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/numbering" Target="numbering.xml"/>',
                        "",
                    )
                    target.writestr(item, relationships.encode("utf-8"))
                continue
            target.writestr(item, source.read(item.filename))
    temporary_path.replace(path)


def test_convert_string_basic():
    """Test basic string conversion."""
    converter = Converter()
    
    md_content = """# Test Heading

This is a paragraph with **bold** and *italic* text.

## Second Heading

Another paragraph."""
    
    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
        output_path = f.name
    
    try:
        converter.convert_string(md_content, output_path)
        
        # Verify file was created
        assert Path(output_path).exists()
        
        # Verify content
        from docx import Document
        doc = Document(output_path)
        assert len(doc.paragraphs) > 0
        
        # Check headings
        assert "Test Heading" in doc.paragraphs[0].text
        assert "Second Heading" in [p.text for p in doc.paragraphs]
    finally:
        # Cleanup
        Path(output_path).unlink(missing_ok=True)


def test_generated_docx_hides_python_docx_fingerprints(tmp_path):
    """Generated files should not expose python-docx defaults to upload platforms."""
    docx_path = tmp_path / "research-report.docx"
    md_content = """# 研究方案

这是一个用于上传审核的正式文档段落。

## 研究目标

确保生成后的 Word 文档不携带默认程序生成指纹。
"""

    converter = Converter(config_override={"metadata": {"author": "md2docx"}})
    converter.convert_string(md_content, str(docx_path))

    core = _docx_xml(docx_path, "docProps/core.xml")
    core_text = ET.tostring(core, encoding="unicode")
    app = _docx_xml(docx_path, "docProps/app.xml")
    app_ns = {"ep": "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"}

    assert "python-docx" not in core_text
    assert "generated by python-docx" not in core_text
    assert "2013-12-23T23:15:00Z" not in core_text

    creator = core.find("{http://purl.org/dc/elements/1.1/}creator")
    title = core.find("{http://purl.org/dc/elements/1.1/}title")
    assert creator is not None
    assert creator.text == "md2docx"
    assert title is not None
    assert title.text == "研究方案"

    words = app.find("ep:Words", app_ns)
    characters = app.find("ep:Characters", app_ns)
    paragraphs = app.find("ep:Paragraphs", app_ns)
    assert words is not None and int(words.text or "0") > 0
    assert characters is not None and int(characters.text or "0") > 0
    assert paragraphs is not None and int(paragraphs.text or "0") > 0


def test_to_document_sets_clean_default_core_properties():
    """Document objects returned by the API should already have clean metadata."""
    doc = Converter().to_document("# 标题\n\n正文。")
    props = doc.core_properties

    assert props.author
    assert props.author != "python-docx"
    assert props.comments == ""
    assert props.title == "标题"
    assert props.created.year != 2013
    assert props.modified.year != 2013


def test_convert_file():
    """Test file conversion."""
    converter = Converter()
    
    # Create temporary Markdown file
    md_content = """# File Test

This is a test from a file."""
    
    with tempfile.NamedTemporaryFile(mode='w', suffix='.md', delete=False) as f:
        f.write(md_content)
        md_path = f.name
    
    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
        docx_path = f.name
    
    try:
        converter.convert(md_path, docx_path)
        
        # Verify output file exists
        assert Path(docx_path).exists()
        
        # Verify content
        from docx import Document
        doc = Document(docx_path)
        assert "File Test" in doc.paragraphs[0].text
    finally:
        # Cleanup
        Path(md_path).unlink(missing_ok=True)
        Path(docx_path).unlink(missing_ok=True)


def test_convert_file_embeds_relative_markdown_image(tmp_path):
    """Relative Markdown image paths should resolve from the Markdown file directory."""
    image_path = tmp_path / "assets" / "png" / "chart.png"
    _write_sample_png(image_path)
    md_path = tmp_path / "report.md"
    docx_path = tmp_path / "report.docx"
    md_path.write_text(
        "# Report\n\nBefore image.\n\n![Chart](assets/png/chart.png)\n\nAfter image.",
        encoding="utf-8",
    )

    Converter().convert(str(md_path), str(docx_path))

    from docx import Document

    doc = Document(docx_path)
    assert len(doc.inline_shapes) == 1
    assert _docx_media_files(docx_path)


def test_convert_file_embeds_relative_image_with_chinese_filename(tmp_path):
    """Percent-encoded image URLs from Mistune should resolve to local Unicode paths."""
    image_path = tmp_path / "assets" / "png" / "图4-1_能力得分比较.png"
    _write_sample_png(image_path)
    md_path = tmp_path / "report.md"
    docx_path = tmp_path / "report.docx"
    md_path.write_text(
        "![图4-1 能力得分比较](assets/png/图4-1_能力得分比较.png)",
        encoding="utf-8",
    )

    Converter().convert(str(md_path), str(docx_path))

    from docx import Document

    doc = Document(docx_path)
    assert len(doc.inline_shapes) == 1
    assert _docx_media_files(docx_path)


def test_convert_string_embeds_absolute_markdown_image(tmp_path):
    """Absolute Markdown image paths should be embedded without a base directory."""
    image_path = tmp_path / "absolute.png"
    docx_path = tmp_path / "absolute.docx"
    _write_sample_png(image_path)

    Converter().convert_string(f"![Chart]({image_path})", str(docx_path))

    from docx import Document

    doc = Document(docx_path)
    assert len(doc.inline_shapes) == 1
    assert _docx_media_files(docx_path)


def test_missing_markdown_image_renders_placeholder(tmp_path):
    """Missing Markdown images should leave a visible placeholder instead of disappearing."""
    md_path = tmp_path / "report.md"
    docx_path = tmp_path / "report.docx"
    md_path.write_text("![Missing](assets/png/missing.png)", encoding="utf-8")

    Converter().convert(str(md_path), str(docx_path))

    from docx import Document

    doc = Document(docx_path)
    assert "[Missing image: assets/png/missing.png]" in [p.text for p in doc.paragraphs]
    assert _docx_media_files(docx_path) == []


def test_unsupported_markdown_image_format_renders_placeholder(tmp_path):
    """Unsupported image formats should be visible in the output document."""
    image_path = tmp_path / "assets" / "graphic.webp"
    image_path.parent.mkdir(parents=True, exist_ok=True)
    image_path.write_bytes(b"webp")
    md_path = tmp_path / "report.md"
    docx_path = tmp_path / "report.docx"
    md_path.write_text("![Graphic](assets/graphic.webp)", encoding="utf-8")

    Converter().convert(str(md_path), str(docx_path))

    from docx import Document

    doc = Document(docx_path)
    assert "[Unsupported image format: assets/graphic.webp]" in [
        p.text for p in doc.paragraphs
    ]
    assert _docx_media_files(docx_path) == []


def test_convert_file_not_found():
    """Test error handling for missing file."""
    converter = Converter()
    
    with pytest.raises(FileNotFoundError):
        converter.convert("nonexistent.md", "output.docx")


def test_to_document():
    """Test to_document method."""
    converter = Converter()
    
    md_content = "# Test\n\nParagraph."
    doc = converter.to_document(md_content)
    
    assert isinstance(doc, DocxDocument)
    assert len(doc.paragraphs) > 0


def test_word_template_preserves_headers_footers_media_and_geometry(tmp_path):
    """A Word template should contribute its complete header/footer package unchanged."""
    image_path = tmp_path / "header-logo.png"
    template_path = tmp_path / "template.docx"
    output_path = tmp_path / "output.docx"
    _write_sample_png(image_path)
    _create_word_template(template_path, image_path)

    converter = Converter(
        config_override={"document": {"page_size": "A4", "margin_top": "3cm"}},
        word_template=str(template_path),
    )
    converter.convert_string("# Generated heading\n\nGenerated body.", str(output_path))

    template = Document(template_path)
    output = Document(output_path)
    assert "Generated body." in [paragraph.text for paragraph in output.paragraphs]
    assert "Template body placeholder" not in [paragraph.text for paragraph in output.paragraphs]
    assert _word_template_parts(template_path) == _word_template_parts(output_path)

    template_section = template.sections[0]
    output_section = output.sections[0]
    for attribute in (
        "page_width",
        "page_height",
        "top_margin",
        "bottom_margin",
        "left_margin",
        "right_margin",
        "header_distance",
        "footer_distance",
        "different_first_page_header_footer",
    ):
        assert getattr(output_section, attribute) == getattr(template_section, attribute)
    assert output.settings.odd_and_even_pages_header_footer is True


def test_word_template_without_numbering_part_still_converts(tmp_path):
    """A header-only template without numbering.xml should remain usable."""
    template_path = tmp_path / "header-only.docx"
    output_path = tmp_path / "output.docx"
    template = Document()
    template.sections[0].header.paragraphs[0].add_run("Template header")
    template.save(template_path)
    _remove_numbering_part(template_path)

    Converter(word_template=str(template_path)).convert_string(
        "# Generated heading\n\nGenerated body.",
        str(output_path),
    )

    output = Document(output_path)
    assert "Generated body." in [paragraph.text for paragraph in output.paragraphs]
    assert output.sections[0].header.paragraphs[0].text == "Template header"


def test_word_template_without_numbering_part_supports_ordered_lists(tmp_path):
    """Missing numbering definitions should be created when a list needs them."""
    template_path = tmp_path / "header-only.docx"
    output_path = tmp_path / "output.docx"
    template = Document()
    template.sections[0].header.paragraphs[0].add_run("Template header")
    template.save(template_path)
    _remove_numbering_part(template_path)

    Converter(word_template=str(template_path)).convert_string(
        "1. First\n2. Second",
        str(output_path),
    )

    output = Document(output_path)
    assert [paragraph.text for paragraph in output.paragraphs] == ["First", "Second"]
    with ZipFile(output_path) as archive:
        assert "word/numbering.xml" in archive.namelist()
        assert b"<w:numPr>" in archive.read("word/document.xml")


def test_word_template_without_numbering_part_supports_bullet_lists(tmp_path):
    """Missing numbering definitions should not block default bullet lists."""
    template_path = tmp_path / "header-only.docx"
    output_path = tmp_path / "output.docx"
    template = Document()
    template.sections[0].header.paragraphs[0].add_run("Template header")
    template.save(template_path)
    _remove_numbering_part(template_path)

    Converter(word_template=str(template_path)).convert_string(
        "* First\n* Second",
        str(output_path),
    )

    output = Document(output_path)
    assert [paragraph.text for paragraph in output.paragraphs] == ["•\tFirst", "•\tSecond"]


def test_word_template_rejects_missing_file(tmp_path):
    with pytest.raises(FileNotFoundError, match="Word template not found"):
        Converter(word_template=str(tmp_path / "missing.docx")).to_document("# Test")


def test_word_template_rejects_non_docx_path(tmp_path):
    template_path = tmp_path / "template.txt"
    template_path.write_text("not a Word document", encoding="utf-8")

    with pytest.raises(ValueError, match="must be a .docx file"):
        Converter(word_template=str(template_path)).to_document("# Test")


def test_word_template_rejects_multiple_sections(tmp_path):
    image_path = tmp_path / "header-logo.png"
    template_path = tmp_path / "multi-section.docx"
    _write_sample_png(image_path)
    _create_word_template(template_path, image_path, sections=2)

    with pytest.raises(ValueError, match="exactly one section"):
        Converter(word_template=str(template_path)).to_document("# Test")


def test_converter_with_template():
    """Test converter with template."""
    converter = Converter(template="default")
    
    md_content = "# Test"
    doc = converter.to_document(md_content)
    
    assert isinstance(doc, DocxDocument)
