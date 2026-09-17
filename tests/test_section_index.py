"""End-to-end checks for optional per-section heading index pages."""
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from md2docx import Converter
from md2docx.config_document import ConfigDocument
from md2docx.config_utils import load_yaml_config


def _converter(**kwargs):
    return Converter(config_override={'document': {'section_index': True}}, **kwargs)


def _text(paragraph):
    return ''.join(paragraph._p.xpath('.//w:t/text()'))


def _links(doc):
    return doc._element.xpath('.//w:body//w:hyperlink')


def _body_texts(doc):
    return [''.join(p.xpath('.//w:t/text()')) for p in doc._element.body.xpath('.//w:p')]


def _entries(table):
    return [p for row in table.rows[1:] for p in row.cells[0].paragraphs]


def test_section_index_is_off_by_default_and_old_configs_keep_the_same_body():
    markdown = '## Section\n\n### Child\n\n#### Detail\n\nText'
    old = Converter(config_data={'document': {'page_size': 'A4'}}).to_document(markdown)
    disabled = Converter(config_data={'document': {'page_size': 'A4', 'section_index': False}}).to_document(markdown)
    assert old._element.body.xml == disabled._element.body.xml
    assert len(old.paragraphs) == 4
    assert not _links(Converter().to_document(markdown))


def test_indexes_follow_source_order_and_stop_at_h1_or_h2():
    markdown = '''### Orphan

## First

### Alpha

#### Detail

##### Too deep

### Beta

## Empty

# New chapter

### Outside

## Last

#### Direct detail
'''
    doc = _converter().to_document(markdown)
    assert _body_texts(doc) == [
        'Orphan', 'First', 'Alpha', 'Detail', 'Beta',
        'First', 'Alpha', 'Detail', 'Too deep', 'Beta',
        'Empty', 'Empty', 'New chapter', 'Outside',
        'Last', 'Direct detail', 'Last', 'Direct detail',
    ]
    assert len(_links(doc)) == 4
    assert [p.text for p in doc.paragraphs if p.style.name == 'Heading 2'] == [
        'First', 'Empty', 'Last',
    ]


def test_empty_section_keeps_a_title_only_index_without_leading_blank_page():
    doc = _converter().to_document('## Empty\n\nBody')
    heading, body = doc.paragraphs
    assert len(doc.tables[0].rows) == 1
    title = doc.tables[0].cell(0, 0).paragraphs[0]
    assert _body_texts(doc) == ['Empty', 'Empty', 'Body']
    assert title.paragraph_format.page_break_before is False
    assert title.paragraph_format.keep_with_next is False
    assert heading.paragraph_format.page_break_before is True
    assert body.paragraph_format.page_break_before is None
    assert not doc._element.xpath('.//w:br[@w:type="page"]')


def test_indexes_have_indentation_clickable_unique_targets_and_no_page_numbers(tmp_path):
    output = tmp_path / 'indexed.docx'
    _converter(config_data={
        'heading3': {'alignment': 'right', 'first_line_indent': 1, 'font_size': '16pt'},
        'heading4': {'alignment': 'right', 'first_line_indent': 2, 'font_size': '12pt'},
    }).convert_string('Intro\n\n## A\n\n### Same\n\n#### Same\n\n## B\n\n### Same', str(output))
    doc = Document(output)
    links = _links(doc)
    anchors = [link.get(qn('w:anchor')) for link in links]
    assert len(set(anchors)) == 3
    starts = doc._element.xpath('.//w:bookmarkStart')
    assert {node.get(qn('w:name')) for node in starts} == set(anchors)
    assert all(node.getparent().xpath('./w:pPr/w:pStyle') for node in starts)
    entries = [p for table in doc.tables for p in _entries(table)]
    assert [p.paragraph_format.left_indent for p in entries] == [Pt(0), Pt(0), Pt(0)]
    assert [p.paragraph_format.first_line_indent for p in entries] == [Pt(16), Pt(24), Pt(16)]
    assert all(p.alignment == WD_ALIGN_PARAGRAPH.RIGHT for p in entries)
    assert all(p.style.name == 'Normal' for p in entries)
    assert all(p._p.xpath('./w:pPr/w:outlineLvl/@w:val') == ['9'] for p in entries)
    assert not doc._element.xpath('.//w:instrText | .//w:fldSimple')
    assert all(table.cell(0, 0).paragraphs[0].paragraph_format.page_break_before is True for table in doc.tables)


def test_indexes_preserve_heading_formatting_and_ignore_fenced_headings():
    markdown = '## Main\n\n### Text **bold**\n\n```text\n## Fake\n### Fake child\n```'
    doc = _converter(config_data={'heading3': {
        'font_name': 'Arial', 'font_size': '13pt', 'bold': False,
    }}).to_document(markdown)
    links = _links(doc)
    assert len(links) == 1
    assert ''.join(links[0].xpath('.//w:t/text()')) == 'Text bold'
    assert links[0].xpath('.//w:rPr/w:b/@w:val') == ['0', '0']
    assert links[0].xpath('.//w:rPr/w:rFonts/@w:ascii') == ['Arial', 'Arial']


def test_chinese_outline_semantic_levels_are_indexed_even_with_custom_styles():
    markdown = '一、总章\n\n（一）本节\n\n1. 事项\n\n（1）细节\n\n二、下一章\n\n1. 节外事项'
    doc = _converter(config_data={'document': {'outline_mode': 'on'}, 'outline': {'level2_style': 'paragraph'}}).to_document(markdown)
    assert _body_texts(doc) == [
        '一、总章', '（一）本节', '1. 事项', '（1）细节',
        '（一）本节', '1. 事项', '（1）细节', '二、下一章', '1. 节外事项',
    ]
    assert len(_links(doc)) == 2


def test_converter_reuse_does_not_leak_heading_targets():
    converter = _converter()
    converter.to_document('## First\n\n### Old child')
    doc = converter.to_document('## Second\n\n#### New child')
    assert _body_texts(doc) == ['Second', 'New child', 'Second', 'New child']
    assert len(_links(doc)) == 1


def test_word_template_geometry_headers_footers_and_bookmarks_are_preserved(tmp_path):
    template = Document()
    section = template.sections[0]
    section.page_width = Cm(18)
    section.header.paragraphs[0].text = 'Header'
    section.footer.paragraphs[0].text = 'Footer'
    bookmark = OxmlElement('w:bookmarkStart')
    bookmark.set(qn('w:id'), '42')
    bookmark.set(qn('w:name'), '_md2docx_index_43')
    section.header.paragraphs[0]._p.append(bookmark)
    end = OxmlElement('w:bookmarkEnd')
    end.set(qn('w:id'), '42')
    section.header.paragraphs[0]._p.append(end)
    template.add_paragraph('Placeholder')
    path = tmp_path / 'template.docx'
    template.save(path)
    doc = _converter(word_template=str(path)).to_document('## Main\n\n### Child')
    assert len(doc.sections) == 1
    assert doc.sections[0].page_width == section.page_width
    assert doc.sections[0].header._element.xml == section.header._element.xml
    assert doc.sections[0].footer._element.xml == section.footer._element.xml
    assert 'Placeholder' not in [_text(p) for p in doc.paragraphs]
    assert _links(doc)[0].get(qn('w:anchor')) == '_md2docx_index_44'


def test_long_index_keeps_all_entries_and_allows_page_flow():
    markdown = '## Main\n\n' + '\n\n'.join(f'### Item {i}' for i in range(100))
    doc = _converter().to_document(markdown)
    assert len(_links(doc)) == 100
    entries = _entries(doc.tables[0])
    assert [_text(p) for p in entries] == [f'Item {i}' for i in range(100)]
    assert all(p.paragraph_format.keep_with_next is False for p in entries)
    assert all(p.paragraph_format.page_break_before is False for p in entries)


def test_index_layout_matches_reference_frame_centering_and_group_spacing():
    doc = _converter(config_data={
        'heading2': {'font_size': '22pt', 'alignment': 'center', 'bold': True, 'space_after': '24pt'},
        'heading3': {'font_size': '16pt', 'alignment': 'center', 'space_before': '20pt'},
        'heading4': {'font_size': '16pt', 'alignment': 'center', 'space_before': '0pt'},
    }).to_document('## Main\n\n### 1.1 Group\n\n#### 1.1.1 Child\n\n### 1.2 Group\n\n#### 1.2.1 Child')
    title = doc.tables[0].cell(0, 0).paragraphs[0]
    assert title.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert all(run.font.size == Pt(22) and run.bold for run in title.runs)
    frame = doc.tables[0]
    assert frame._tbl.xpath('./w:tblPr/w:tblBorders/*[not(starts-with(local-name(), "inside"))]/@w:val') == ['double'] * 4
    assert not frame._tbl.xpath('./w:tr/w:trPr/w:trHeight')
    assert frame._tbl.xpath('./w:tr/w:tc/w:tcPr/w:vAlign/@w:val') == ['center'] * 5
    paragraphs = _entries(frame)
    assert [p.paragraph_format.space_before for p in paragraphs] == [Pt(20), Pt(0), Pt(20), Pt(0)]
    assert [p.paragraph_format.keep_with_next for p in paragraphs] == [True, False, True, False]
    assert all(p.alignment == WD_ALIGN_PARAGRAPH.CENTER for p in paragraphs)


@pytest.mark.parametrize('entry_count', [1, 6, 19, 100])
def test_index_frame_fits_content_and_starts_with_its_title(tmp_path, entry_count):
    """No page-sized row or fixed width may push the frame past its title."""
    template = Document()
    template.sections[0].page_width = Cm(18)
    template.sections[0].page_height = Cm(20)
    template.sections[0].header.paragraphs[0].text = 'Template header'
    template.sections[0].footer.paragraphs[0].text = 'Template footer'
    template_path = tmp_path / 'template.docx'
    template.save(template_path)
    title_text = 'Long section title with enough text to wrap across several lines'
    markdown = f'## {title_text}\n\n' + '\n\n'.join(
        f'### Entry {i}\n\nBody text.' for i in range(entry_count)
    )
    result_path = tmp_path / 'result.docx'
    _converter(
        word_template=str(template_path),
        config_data={'heading2': {
            'font_size': '28pt', 'space_before': '12pt', 'space_after': '24pt',
        }},
    ).convert_string(markdown, str(result_path))
    doc = Document(result_path)
    frame = doc.tables[0]
    title = frame.cell(0, 0).paragraphs[0]
    body_title = doc.paragraphs[0]
    assert title.text == body_title.text == title_text
    assert len(frame.rows) == entry_count + 1
    assert all(len(row.cells[0].paragraphs) == 1 for row in frame.rows)
    assert frame._tbl.getnext() is body_title._p
    assert not doc._element.body.xpath('./w:p/w:pPr/w:outlineLvl[@w:val="9"]')
    assert title.paragraph_format.keep_with_next is True
    assert title.paragraph_format.page_break_before is False
    assert title.paragraph_format.space_after == Pt(24)
    assert body_title.paragraph_format.page_break_before is True
    assert frame.autofit is True
    assert frame._tbl.xpath('./w:tblPr/w:tblW/@w:type') == ['auto']
    assert frame._tbl.xpath('./w:tr/w:tc/w:tcPr/w:tcW/@w:type') == ['auto'] * (entry_count + 1)
    assert not frame._tbl.xpath('./w:tr/w:trPr/w:trHeight')
    assert len(frame._tbl.xpath('./w:tr/w:trPr/w:cantSplit')) == entry_count + 1
    # Padding belongs to the outside of the frame, not every entry row.
    assert frame._tbl.xpath('./w:tr/w:tc/w:tcPr/w:tcMar/w:top/@w:w') == ['200'] + ['0'] * entry_count
    assert frame._tbl.xpath('./w:tr/w:tc/w:tcPr/w:tcMar/w:bottom/@w:w') == ['0'] * entry_count + ['200']
    assert frame._tbl.xpath('./w:tblPr/w:tblBorders/w:insideH/@w:val') == ['nil']
    entries = _entries(frame)
    assert [_text(p) for p in entries] == [f'Entry {i}' for i in range(entry_count)]
    assert all(p.paragraph_format.page_break_before is False for p in entries)
    assert doc.sections[0]._sectPr.xml == template.sections[0]._sectPr.xml
    assert doc.sections[0].header._element.xml == template.sections[0].header._element.xml
    assert doc.sections[0].footer._element.xml == template.sections[0].footer._element.xml


def test_old_gui_config_gets_switch_and_saved_toggle_controls_conversion(tmp_path):
    defaults = load_yaml_config(Path('md2docx/templates/default.yaml'))
    path = tmp_path / 'config.yaml'
    path.write_text('document:\n  page_size: A4\n', encoding='utf-8')
    config = ConfigDocument.load(path, defaults)
    assert config.draft_config['document']['section_index'] is False
    config.draft_config['document']['section_index'] = True
    config.update_draft(config.draft_config)
    config.save()
    saved = ConfigDocument.load(path, defaults).saved_config
    doc = Converter(config_data=saved).to_document('## A\n\n### B')
    assert len(_links(doc)) == 1


@pytest.mark.parametrize('outline', [False, True])
def test_index_uses_saved_heading_styles_by_semantic_level(tmp_path, outline):
    """Distinct H2/H3/H4 settings survive YAML save/reload and DOCX round trips."""
    config = ConfigDocument.from_defaults(load_yaml_config(Path('md2docx/templates/default.yaml')))
    config.draft_config['document'].update(section_index=True, outline_mode='on' if outline else 'off')
    alignments = ('center', 'right', 'left')
    fonts = ('Arial', '宋体', '楷体')
    sizes = (22, 15.5, 10.5)
    for level, font, size, alignment in zip((2, 3, 4), fonts, sizes, alignments):
        config.draft_config[f'heading{level}'].update(
            font_name=font, font_size=f'{size}pt', font_color='#123456',
            bold=level != 3, italic=level == 4, alignment=alignment,
            first_line_indent=level - 2, space_before=f'{level}.5pt',
            space_after=f'{level + 1}.5pt', line_spacing=1.2,
        )
    config.update_draft(config.draft_config)
    path = tmp_path / 'styles.yaml'
    config.save_as(path)
    saved = ConfigDocument.load(path, {}).saved_config
    assert saved == config.saved_config

    markdown = '（一）本节\n\n1. 事项\n\n（1）细节' if outline else '## Main\n\n### Child\n\n#### Detail'
    output = tmp_path / 'styled.docx'
    Converter(config_data=saved).convert_string(markdown, str(output))
    doc = Document(output)
    indexed = [doc.tables[0].cell(0, 0).paragraphs[0], *_entries(doc.tables[0])]
    for level, paragraph, font, size, alignment in zip((2, 3, 4), indexed, fonts, sizes, alignments):
        assert paragraph._p.xpath('.//w:rPr/w:rFonts/@w:eastAsia') == [font]
        assert paragraph._p.xpath('.//w:rPr/w:sz/@w:val') == [str(int(size * 2))]
        assert paragraph._p.xpath('.//w:rPr/w:color/@w:val') == ['123456']
        assert paragraph._p.xpath('.//w:pPr/w:jc/@w:val') == [alignment]
        formatting = paragraph.paragraph_format
        assert formatting.first_line_indent == Pt((level - 2) * size)
        assert formatting.space_before == Pt(level + 0.5)
        assert formatting.space_after == Pt(level + 1.5)
        assert formatting.line_spacing == 1.2
        assert paragraph._p.xpath('./w:pPr/w:outlineLvl/@w:val') == ['9']
    # Template/source styles are not modified when building the index.
    if not outline:
        for source, index in zip(doc.paragraphs, indexed):
            assert source.paragraph_format.first_line_indent == index.paragraph_format.first_line_indent
            assert source.paragraph_format.line_spacing == index.paragraph_format.line_spacing


def test_index_preserves_template_first_even_footers_and_page_number_fields(tmp_path):
    template = Document()
    section = template.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(1.9)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.3)
    section.different_first_page_header_footer = True
    template.settings.odd_and_even_pages_header_footer = True
    for kind in ('header', 'footer', 'first_page_header', 'first_page_footer', 'even_page_header', 'even_page_footer'):
        paragraph = getattr(section, kind).paragraphs[0]
        paragraph.text = f'Template {kind}'
        if 'footer' in kind:
            field = OxmlElement('w:fldSimple')
            field.set(qn('w:instr'), ' PAGE ')
            paragraph._p.append(field)
    numbering = OxmlElement('w:pgNumType')
    numbering.set(qn('w:start'), '7')
    numbering.set(qn('w:fmt'), 'upperRoman')
    section._sectPr.append(numbering)
    path = tmp_path / 'template.docx'
    output = tmp_path / 'result.docx'
    template.save(path)
    _converter(word_template=str(path)).convert_string('## First\n\n### Child\n\n## Empty', str(output))
    result = Document(output)
    assert len(result.sections) == 1
    assert result.sections[0]._sectPr.xml == section._sectPr.xml
    assert result.settings.odd_and_even_pages_header_footer
    for kind in ('header', 'footer', 'first_page_header', 'first_page_footer', 'even_page_header', 'even_page_footer'):
        assert getattr(result.sections[0], kind)._element.xml == getattr(section, kind)._element.xml
    assert not result._element.body.xpath('.//w:fldSimple | .//w:instrText')
    assert result.tables[0].columns[0].width == section.page_width - section.left_margin - section.right_margin


@pytest.mark.parametrize('outline_mode', ['auto', 'on'])
@pytest.mark.parametrize('use_template', [False, True])
def test_markdown_sections_do_not_turn_body_outline_titles_into_index_pages(
    tmp_path, outline_mode, use_template,
):
    """Chinese subheadings inside Markdown chapters must not split body pages."""
    markdown = '''# 技术方案

## 1. 项目理解

### 1.1 项目背景

#### 1.1.1 运行背景

一、政策背景

（一）综合帮扶

综合帮扶的正文。

（二）主动发现

主动发现的正文。

二、现实基础

（一）统一平台

统一平台的正文。

#### 1.1.2 服务需求

服务需求的正文。

### 1.2 项目目标

目标正文。

## 2. 服务安排

### 2.1 实施计划

计划正文。
'''
    kwargs = {'config_data': {'document': {'outline_mode': outline_mode}}}
    if use_template:
        template = Document()
        template.sections[0].header.paragraphs[0].text = '保留模板页眉'
        template.sections[0].footer.paragraphs[0].text = '保留模板页脚'
        template_path = tmp_path / 'template.docx'
        template.save(template_path)
        kwargs['word_template'] = str(template_path)
    baseline = Converter(**kwargs).to_document(markdown)
    output = tmp_path / 'result.docx'
    _converter(**kwargs).convert_string(markdown, str(output))
    result = Document(output)
    index_titles = [table.cell(0, 0).paragraphs[0] for table in result.tables]
    assert [p.text for p in index_titles] == ['1. 项目理解', '2. 服务安排']
    assert len(result.tables) == 2
    assert [_text(p) for p in _entries(result.tables[0])] == [
        '1.1 项目背景', '1.1.1 运行背景', '1.1.2 服务需求', '1.2 项目目标',
    ]
    assert [_text(p) for p in _entries(result.tables[1])] == ['2.1 实施计划']
    body = [p for p in result.paragraphs if not p._p.xpath('./w:pPr/w:outlineLvl[@w:val="9"]')]
    assert [p.text for p in body] == [p.text for p in baseline.paragraphs]
    # Index boundaries may start source H2 on a page. Every other source
    # paragraph keeps its original pagination and formatting, including outlines.
    for actual, expected in zip(body, baseline.paragraphs):
        if actual.style.name == 'Heading 2':
            assert actual.paragraph_format.page_break_before is True
        else:
            actual_format = actual._p.pPr.xml if actual._p.pPr is not None else None
            expected_format = expected._p.pPr.xml if expected._p.pPr is not None else None
            assert actual_format == expected_format
    assert result.sections[0]._sectPr.xml == baseline.sections[0]._sectPr.xml
    assert result.sections[0].header._element.xml == baseline.sections[0].header._element.xml
    assert result.sections[0].footer._element.xml == baseline.sections[0].footer._element.xml


def test_chinese_outline_index_with_document_title_and_converter_reuse():
    """A top-level document title still allows a pure Chinese chapter outline."""
    converter = _converter(config_data={'document': {'outline_mode': 'on'}})
    converter.to_document('## Markdown section\n\n### Child')
    result = converter.to_document('# 文档标题\n\n一、章节\n\n（一）分节\n\n1. 事项\n\n（1）细节')
    assert _body_texts(result) == [
        '文档标题', '一、章节', '（一）分节', '1. 事项', '（1）细节',
        '（一）分节', '1. 事项', '（1）细节',
    ]
