"""Tests for selectable DOCX table layout presets."""
from pathlib import Path

from docx.oxml.ns import qn

from md2docx import Converter
from md2docx.config_utils import save_yaml_config


TABLE_MARKDOWN = """
| 姓名 | 分数 |
|------|------|
| 张三 | 95 |
| 李四 | 88 |
"""


def _render_table(layout=None, **table_overrides):
    table_config = dict(table_overrides)
    if layout is not None:
        table_config["layout"] = layout
    document = Converter(config_override={"table": table_config}).to_document(
        TABLE_MARKDOWN
    )
    return document.tables[0]


def _table_border(table, side):
    borders = table._tbl.tblPr.first_child_found_in("w:tblBorders")
    return None if borders is None else borders.find(qn(f"w:{side}"))


def _cell_border(cell, side):
    borders = cell._tc.tcPr.first_child_found_in("w:tcBorders")
    return None if borders is None else borders.find(qn(f"w:{side}"))


def _border_attributes(border):
    assert border is not None
    return {
        "value": border.get(qn("w:val")),
        "size": border.get(qn("w:sz")),
        "color": border.get(qn("w:color")),
    }


def _cell_shading(cell):
    shading = cell._tc.tcPr.find(qn("w:shd"))
    return None if shading is None else shading.get(qn("w:fill"))


def test_missing_and_invalid_layout_preserve_theme_grid():
    missing_converter = Converter()
    missing_converter.style_manager.config["table"].pop("layout", None)
    missing = missing_converter.to_document(TABLE_MARKDOWN).tables[0]
    invalid = _render_table("unsupported-layout")

    assert missing.style.name == "Light Grid Accent 1"
    assert invalid.style.name == "Light Grid Accent 1"


def test_explicit_accent_grid_preserves_style_and_configured_shading():
    table = _render_table(
        "accent_grid",
        alternating_rows=True,
        row_background_odd="#FFFFFF",
        row_background_even="#F9F9F9",
    )

    assert table.style.name == "Light Grid Accent 1"
    assert _cell_shading(table.rows[0].cells[0]) == "F2F2F2"
    assert _cell_shading(table.rows[1].cells[0]) == "FFFFFF"
    assert _cell_shading(table.rows[2].cells[0]) == "F9F9F9"


def test_plain_grid_has_black_half_point_grid_and_gray_header():
    table = _render_table(
        "plain_grid",
        alternating_rows=True,
        row_background_odd="#FF0000",
        row_background_even="#00FF00",
    )

    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        assert _border_attributes(_table_border(table, side)) == {
            "value": "single",
            "size": "4",
            "color": "000000",
        }
    assert table._tbl.tblPr.find(qn("w:tblStyle")) is None
    assert _cell_shading(table.rows[0].cells[0]) == "F2F2F2"
    assert _cell_shading(table.rows[1].cells[0]) is None
    assert _cell_shading(table.rows[2].cells[0]) is None


def test_three_line_has_only_top_header_separator_and_bottom_borders():
    table = _render_table(
        "three_line",
        alternating_rows=True,
        row_background_odd="#FF0000",
        row_background_even="#00FF00",
    )

    for side in ("top", "bottom"):
        assert _border_attributes(_table_border(table, side)) == {
            "value": "single",
            "size": "12",
            "color": "000000",
        }
    for side in ("left", "right", "insideH", "insideV"):
        assert _border_attributes(_table_border(table, side))["value"] == "nil"
    for cell in table.rows[0].cells:
        assert _border_attributes(_cell_border(cell, "bottom")) == {
            "value": "single",
            "size": "6",
            "color": "000000",
        }
    for row in table.rows[1:]:
        for cell in row.cells:
            assert _cell_border(cell, "bottom") is None
    assert table._tbl.tblPr.find(qn("w:tblStyle")) is None
    assert all(_cell_shading(cell) is None for row in table.rows for cell in row.cells)


def test_saved_yaml_layout_reaches_the_renderer(tmp_path: Path):
    config_path = tmp_path / "three-line.yaml"
    save_yaml_config(
        config_path,
        {
            "table": {
                "layout": "three_line",
                "font_name": "仿宋",
                "font_size": "12pt",
            }
        },
    )

    table = Converter(style_config=str(config_path)).to_document(TABLE_MARKDOWN).tables[0]

    assert _border_attributes(_table_border(table, "top"))["size"] == "12"
    assert (
        _border_attributes(_cell_border(table.rows[0].cells[0], "bottom"))["size"]
        == "6"
    )
