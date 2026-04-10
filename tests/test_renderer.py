"""
Tests for renderer module.
"""
import pytest
from docx import Document
from md2docx.renderer import DocxRenderer
from md2docx.styles import StyleManager


def test_renderer_initialization():
    """Test renderer can be initialized."""
    doc = Document()
    style_manager = StyleManager()
    renderer = DocxRenderer(doc, style_manager)
    
    assert renderer is not None
    assert renderer.doc is doc
    assert renderer.styles is style_manager


def test_heading_rendering():
    """Test heading rendering with token-based API."""
    doc = Document()
    style_manager = StyleManager()
    renderer = DocxRenderer(doc, style_manager)
    
    # Create token as mistune would
    token = {
        'type': 'heading',
        'attrs': {'level': 1},
        'children': [{'type': 'text', 'raw': 'Level 1 Heading'}]
    }
    
    # Use a mock state
    class MockState:
        pass
    
    renderer.heading(token, MockState())
    
    assert len(doc.paragraphs) == 1
    assert doc.paragraphs[0].text == "Level 1 Heading"


def test_parse_font_size():
    """Test font size parsing."""
    doc = Document()
    style_manager = StyleManager()
    renderer = DocxRenderer(doc, style_manager)
    
    assert renderer._parse_font_size("12pt") == 12
    assert renderer._parse_font_size("18pt") == 18
    assert renderer._parse_font_size("14") == 14
    assert renderer._parse_font_size(16) == 16


def test_parse_color():
    """Test color parsing."""
    doc = Document()
    style_manager = StyleManager()
    renderer = DocxRenderer(doc, style_manager)
    
    color = renderer._parse_color("#FF0000")
    # RGBColor is a tuple-like object
    assert isinstance(color,tuple)
    assert len(color) == 3
    assert color == (255, 0, 0)
    
    color2 = renderer._parse_color("0000FF")
    assert color2 == (0, 0, 255)


def test_render_children():
    """Test render_children method."""
    doc = Document()
    style_manager = StyleManager()
    renderer = DocxRenderer(doc, style_manager)
    
    token = {
        'children': [
            {'type': 'text', 'raw': 'Hello '},
            {'type': 'text', 'raw': 'World'}
        ]
    }
    
    class MockState:
        pass
    
    result = renderer.render_children(token, MockState())
    assert result == ['Hello ', 'World']


def test_thematic_break_is_ignored_by_default():
    """Markdown horizontal rules should be skipped by default."""
    doc = Document()
    style_manager = StyleManager()
    renderer = DocxRenderer(doc, style_manager)

    class MockState:
        pass

    renderer.thematic_break({"type": "thematic_break"}, MockState())

    assert len(doc.paragraphs) == 0


def test_thematic_break_can_be_rendered_when_enabled():
    """Users can opt back into thematic break output via config."""
    doc = Document()
    style_manager = StyleManager()
    style_manager.config.setdefault("document", {})["ignore_thematic_breaks"] = False
    renderer = DocxRenderer(doc, style_manager)

    class MockState:
        pass

    renderer.thematic_break({"type": "thematic_break"}, MockState())

    assert len(doc.paragraphs) == 1
    assert doc.paragraphs[0].text == "_" * 50
