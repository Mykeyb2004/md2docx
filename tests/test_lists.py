"""
Test list rendering functionality.
"""
import tempfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import List, Optional, Tuple
from zipfile import ZipFile

from docx import Document

from md2docx import Converter


WORD_NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}


def _read_paragraph_numbering(docx_path: str) -> List[Tuple[str, Optional[str]]]:
    """Return each paragraph's text with its explicit numbering id if present."""
    with ZipFile(docx_path) as archive:
        document_xml = archive.read('word/document.xml')

    root = ET.fromstring(document_xml)
    paragraphs: List[Tuple[str, Optional[str]]] = []

    for paragraph in root.findall('.//w:body/w:p', WORD_NS):
        text = ''.join(node.text or '' for node in paragraph.findall('.//w:t', WORD_NS)).strip()
        num_id = paragraph.find('./w:pPr/w:numPr/w:numId', WORD_NS)
        paragraphs.append(
            (
                text,
                num_id.get(f'{{{WORD_NS["w"]}}}val') if num_id is not None else None,
            )
        )

    return paragraphs


def test_unordered_list():
    """Test unordered list conversion."""
    converter = Converter()

    md_content = """# Lists Test

* Item 1
* Item 2
* Item 3
"""

    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
        output_path = f.name

    try:
        converter.convert_string(md_content, output_path)
        assert Path(output_path).exists()

        doc = Document(output_path)
        # Should have heading + 3 list items
        assert len(doc.paragraphs) >= 4

        # Check list items exist
        list_items = [p for p in doc.paragraphs if 'List' in p.style.name]
        assert len(list_items) >= 3
    finally:
        Path(output_path).unlink(missing_ok=True)


def test_ordered_list():
    """Test ordered list conversion."""
    converter = Converter()

    md_content = """# Ordered List

1. First item
2. Second item
3. Third item
"""

    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
        output_path = f.name

    try:
        converter.convert_string(md_content, output_path)
        assert Path(output_path).exists()

        doc = Document(output_path)
        list_items = [p for p in doc.paragraphs if 'List' in p.style.name]
        assert len(list_items) >= 3
    finally:
        Path(output_path).unlink(missing_ok=True)


def test_mixed_lists():
    """Test document with both ordered and unordered lists."""
    converter = Converter()

    md_content = """# Mixed Lists

Unordered:
* Apple
* Banana

Ordered:
1. First
2. Second
"""

    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
        output_path = f.name

    try:
        converter.convert_string(md_content, output_path)
        assert Path(output_path).exists()

        doc = Document(output_path)
        list_items = [p for p in doc.paragraphs if 'List' in p.style.name]
        assert len(list_items) >= 4
    finally:
        Path(output_path).unlink(missing_ok=True)


def test_real_document_with_lists():
    """Test converting a more complete document with lists."""
    converter = Converter()

    md_content = """# 项目文档

## 功能列表

项目具有以下功能：

1. Markdown 转换
2. 样式定制
3. 多种格式支持

## 主要特性

- 快速转换
- 高质量输出
- 易于使用
"""

    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
        output_path = f.name

    try:
        converter.convert_string(md_content, output_path)
        assert Path(output_path).exists()

        doc = Document(output_path)
        assert len(doc.paragraphs) > 5
    finally:
        Path(output_path).unlink(missing_ok=True)


def test_ordered_lists_restart_numbering_between_blocks():
    """Each Markdown ordered-list block should restart numbering from 1."""
    converter = Converter()

    md_content = """# 编号重置

（一）第一部分

1. 甲事项
2. 乙事项

（二）第二部分

1. 丙事项
2. 丁事项
"""

    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
        output_path = f.name

    try:
        converter.convert_string(md_content, output_path)
        numbered_paragraphs = {
            text: num_id
            for text, num_id in _read_paragraph_numbering(output_path)
            if text in {'甲事项', '乙事项', '丙事项', '丁事项'}
        }

        assert numbered_paragraphs['甲事项'] is not None
        assert numbered_paragraphs['乙事项'] == numbered_paragraphs['甲事项']
        assert numbered_paragraphs['丙事项'] is not None
        assert numbered_paragraphs['丁事项'] == numbered_paragraphs['丙事项']
        assert numbered_paragraphs['丙事项'] != numbered_paragraphs['甲事项']
    finally:
        Path(output_path).unlink(missing_ok=True)


def test_loose_ordered_lists_keep_word_numbering():
    """Loose ordered lists should still render as numbered Word paragraphs."""
    converter = Converter()

    md_content = """# 松散编号列表

1. 第一项

2. 第二项

3. 第三项
"""

    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
        output_path = f.name

    try:
        converter.convert_string(md_content, output_path)
        numbering = {
            text: num_id for text, num_id in _read_paragraph_numbering(output_path)
        }

        assert numbering['第一项'] is not None
        assert numbering['第二项'] == numbering['第一项']
        assert numbering['第三项'] == numbering['第一项']
    finally:
        Path(output_path).unlink(missing_ok=True)


def test_ordered_lists_can_render_as_literal_text():
    """Ordered lists can be emitted as plain text markers instead of Word numbering."""
    md_content = """# 文本编号

1. 第一项

2. 第二项

3. 第三项
"""

    with tempfile.NamedTemporaryFile(mode='w', suffix='.yaml', delete=False, encoding='utf-8') as f:
        style_path = f.name
        f.write(
            "list:\n"
            "  ordered_list_as_text: true\n"
            "  number_format: \"1.\"\n"
        )

    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
        output_path = f.name

    try:
        converter = Converter(style_config=style_path)
        converter.convert_string(md_content, output_path)

        doc = Document(output_path)
        texts = [p.text for p in doc.paragraphs]
        assert '1. 第一项' in texts
        assert '2. 第二项' in texts
        assert '3. 第三项' in texts

        numbering = {
            text: num_id for text, num_id in _read_paragraph_numbering(output_path)
        }
        assert numbering['1. 第一项'] is None
        assert numbering['2. 第二项'] is None
        assert numbering['3. 第三项'] is None
    finally:
        Path(output_path).unlink(missing_ok=True)
        Path(style_path).unlink(missing_ok=True)
